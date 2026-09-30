"""ИИ-22, этап 1: презентация PPTX из сохранённого ответа ИИ.

Модели здесь нет и не нужно: построитель детерминированный. Проверяется состав слайдов,
что числа на графиках и в таблицах — ровно из результатов ответа, что SQL в файл не
попадает, лимит 15 слайдов, отказ для ответов без результата, шаблон .potx и эндпоинт
(свой ответ — файл и запись в журнале выгрузок; чужой — 404; план до выполнения — 409).
"""
import copy
import io
import os
import pathlib
import tempfile
import unittest
import zipfile
from datetime import datetime

from fastapi import FastAPI
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches

from backend.ai import api as ai_api
from backend.ai import dialogs as store
from backend.ai import journal
from backend.ai import slides
from backend.ai.export import MSK

SQL = ("SELECT f.account_date AS \"Дата\", SUM(f.sum_receipt_netto_ntu) AS \"Выручка НТУ, ₽\" "
       "FROM dm.data_for_ai_analytic_part_1 AS f WHERE f.account_date >= DATE '2026-09-01' GROUP BY 1")
SRC = {"id": "r1", "title": "выручка НТУ по месяцам", "kind": "sql", "rows": 6}
KPI = {"id": "c1", "type": "kpi", "title": "Выручка НТУ к плану", "source": "r2", "unit": "", "sourceKind": "sql",
       "sourceTitle": "факт и план НТУ",
       "cards": [{"label": "Выручка НТУ, ₽", "value": 10812400, "base": 11250000, "baseLabel": "План, ₽",
                  "delta": -437600.0, "deltaPct": -3.9}]}
LINE = {"id": "c2", "type": "compare", "title": "Выручка НТУ по дням: сентябрь к сентябрю", "source": "r1", "unit": "₽",
        "sourceTitle": "выручка по дням", "sourceKind": "sql", "period": "1–3 сен 2026 и 1–3 сен 2025",
        "x": ["01.09", "02.09", "03.09"], "xTitle": "Дата",
        "series": [{"name": "1–3 сен 2026", "values": [433380.5, 451402.0, 441400.0]},
                   {"name": "1–3 сен 2025", "values": [416977.0, None, 427737.0]}],
        "periods": [{"label": "1–3 сен 2026", "points": 3, "total": 1326182.5},
                    {"label": "1–3 сен 2025", "points": 3, "total": 844714.0}],
        "aggregate": "sum", "changePct": 57.0}
BAR = {"id": "c3", "type": "bar", "title": "Изменение выручки к прошлому году, %", "source": "p1", "unit": "%",
       "sourceTitle": "изменение по ОНПО", "sourceKind": "python", "x": ["Север", "Юг", "Центр"], "xTitle": "ОНПО",
       "series": [{"name": "Изменение, %", "values": [4.2, 2.8, -1.6]}]}
WATERFALL = {"id": "c4", "type": "waterfall", "title": "Факторы изменения выручки, тыс. ₽", "source": "p2",
             "unit": "тыс. ₽", "sourceKind": "python", "sourceTitle": "факторы",
             "x": ["Чеки", "Средний чек"], "series": [{"name": "Вклад", "values": [420.5, -130.2]}],
             "start": 12030.0, "total": 12320.3}
SCATTER = {"id": "c5", "type": "scatter", "title": "Конверсия и средний чек", "source": "r2", "unit": "%",
           "sourceKind": "sql", "xTitle": "Средний чек, ₽", "yTitle": "Конверсия, %",
           "points": [{"x": 380.0, "y": 18.2, "label": "58-123"}, {"x": 410.5, "y": 27.0, "label": "58-140"}]}
AGENT = {
    "ok": True, "question": "Сравни выручку НТУ за сентябрь с прошлым годом", "scopeLabel": "ОНПО «Юг» — 312 АЗС (по ОХД)",
    "depth": "deep", "taskType": "compare", "sql": SQL,
    "analysis": {
        "headline": "Выручка НТУ за 1–3 сентября — 1,33 млн ₽, план месяца пока отстаёт на 3,9 %.",
        "happened": ["Рост к прошлому году по одинаковому числу дней.", "План месяца отстаёт на 3,9 %."],
        "why": ["Возможно, рост дают чеки."], "where": ["Снижение — в ОНПО «Центр»."],
        "limitations": ["Данные пилотной витрины."],
        "claims": {"happened": [{"type": "fact", "sources": [SRC]},
                                {"type": "calc", "sources": [SRC], "formula": ["(10 812 400 / 11 250 000 − 1) × 100 % = −3,89 %"]}],
                   "why": [{"type": "hypothesis", "sources": [], "check": "Чеки по неделям."}],
                   "where": [{"type": "fact", "sources": [SRC]}]},
        "recommendations": [{"action": "Можно разобрать выкладку НТУ на АЗС 58-123.", "basis": "Конверсия 18,2 %.",
                             "effect": "Рост конверсии.", "limits": "Не при ремонте.", "confidence": "средняя",
                             "source": "Анализ витрины (r2)"}],
        "recommendationNote": "Рекомендации носят справочный характер; решение принимает руководитель."},
    "plan": {"period": "1–3 сентября 2026 и 2025", "filters": ["все АЗС общества"]},
    "charts": [KPI, LINE, BAR, WATERFALL, SCATTER],
    "tables": [{"id": "p1", "title": "изменение по ОНПО", "columns": ["ОНПО", "Изменение, %"],
                "rows": [["Север", 4.2], ["Юг", 2.8], ["Центр", -1.6]], "truncated": False}],
    "steps": [{"key": "s1", "kind": "sql", "resultId": "r1", "rows": 6, "purpose": "выручка по дням", "sql": SQL},
              {"key": "s2", "kind": "python", "resultId": "p1", "rows": 3, "purpose": "изменение по ОНПО"}],
    "frame": {"mainResult": "r1"},
    "columns": ["Дата", "Выручка НТУ, ₽"], "rows": [["2026-09-01", 433380.5], ["2026-09-02", 451402]],
    "truncated": False, "notes": [], "context": {"files": [{"id": 1, "name": "План сентябрь.xlsx", "kind": "xlsx"}]},
}
FAST = {"ok": True, "question": "Топ ОНПО по выручке НТУ", "scopeLabel": "вся сеть", "depth": "fast", "taskType": "lookup",
        "summary": "Больше всего выручки НТУ — в ОНПО «Север»: 12,4 млн ₽.", "sql": SQL,
        "columns": ["ОНПО", "Выручка НТУ, ₽"], "rows": [["Север", 12400000], ["Юг", 10950000]],
        "charts": [{"id": "c1", "type": "bar", "title": "Выручка НТУ, ₽", "source": "r1", "unit": "₽", "auto": True,
                    "sourceKind": "sql", "x": ["Север", "Юг"], "series": [{"name": "Выручка НТУ, ₽", "values": [12400000.0, 10950000.0]}]}],
        "tables": [], "steps": [], "frame": {},
        "notes": ["Область данных: вся сеть, фильтр не требуется", "Ограничение вывода: 200 строк"]}
META = {"role": "АУП общества (Юг)", "source": "Витрина ОХД (DWH ЛИКАРД)", "tables": ["dm.data_for_ai_analytic_part_1"],
        "catalog": "встроенный каталог стенда", "now": datetime(2026, 9, 25, 14, 5, tzinfo=MSK), "email": "user@example.com"}


def message(answer, mid=11, journal_id=501):
    return {"id": mid, "journalId": journal_id, "question": answer["question"], "createdAt": 1790330000, "answer": answer}


def deck(answer, **kwargs):
    blob, count = slides.build(message(answer), META, **kwargs)
    prs = Presentation(io.BytesIO(blob))
    assert len(prs.slides) == count
    return blob, prs


def texts(slide) -> str:
    out = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            out.append(shape.text_frame.text)
        if shape.has_table:
            out += [cell.text for row in shape.table.rows for cell in row.cells]
    return "\n".join(out)


def charts(prs):
    return [shape.chart for slide in prs.slides for shape in slide.shapes if shape.has_chart]


class OutlineTests(unittest.TestCase):
    def test_order_title_first_sources_last(self):
        kinds = [s["kind"] for s in slides.outline(message(AGENT), META)]
        self.assertEqual(kinds[0], "title")
        self.assertEqual(kinds[1], "headline")
        self.assertEqual(kinds[-1], "sources")
        self.assertEqual(kinds[-2], "limits")
        self.assertIn("recs", kinds)
        self.assertEqual(kinds.count("chart"), 4)            # KPI — на «Главном», остальные четыре — графиками
        self.assertLessEqual(len(kinds), slides.MAX_SLIDES)

    def test_points_carry_claim_types(self):
        spec = slides.outline(message(AGENT), META)
        head = next(s for s in spec if s["kind"] == "headline")
        points = head["points"] + [p for s in spec if s["kind"] == "points" for p in s["points"]]
        self.assertEqual([p["type"] for p in points], ["fact", "calc", "hypothesis", "fact"])
        calc = points[1]
        self.assertIn("11 250 000", calc["detail"])
        self.assertTrue(points[2]["detail"].startswith("Проверить:"))

    def test_refuses_without_result(self):
        for answer in ({"ok": False, "error": "нет"}, dict(FAST, planCard={"status": "draft"}),
                       {"ok": True, "question": "q", "rows": [], "charts": []}):
            with self.assertRaises(slides.NotPresentable):
                slides.outline(message(dict(answer, question="q")), META)

    def test_slide_cap_names_what_was_left_out(self):
        many = copy.deepcopy(AGENT)
        many["charts"] = [KPI] + [dict(BAR, id=f"b{i}") for i in range(20)]
        spec = slides.outline(message(many), META)
        self.assertEqual(len(spec), slides.MAX_SLIDES)
        limits = next(s for s in spec if s["kind"] == "limits")["items"]
        self.assertTrue(any("не больше 15 слайдов" in item for item in limits))

    def test_fast_answer_puts_headline_on_chart(self):
        spec = slides.outline(message(FAST), META)
        self.assertEqual([s["kind"] for s in spec], ["title", "chart", "sources"])
        self.assertEqual(spec[1]["headline"], FAST["summary"])   # таблица из двух столбцов — это тот же график
        # служебные пометки валидатора (область, лимит строк) — не «ограничения» для слайда


class DeckTests(unittest.TestCase):
    def test_file_is_16x9_with_footer_and_grif(self):
        _blob, prs = deck(AGENT)
        self.assertEqual(prs.slide_width, Inches(13.333))
        self.assertEqual(prs.slide_height, Inches(7.5))
        for index, slide in enumerate(prs.slides, start=1):
            text = texts(slide)
            self.assertIn("LUKOIL ОНПО · Инструмент АУП · Конфиденциально", text)
            self.assertIn(f"{index} / {len(prs.slides)}", text)
        self.assertEqual(prs.core_properties.subject, "Конфиденциально")

    def test_chart_numbers_come_from_results(self):
        _blob, prs = deck(AGENT)
        found = charts(prs)
        self.assertEqual(len(found), 4)
        line = found[0]
        self.assertEqual(list(line.plots[0].categories), LINE["x"])
        self.assertEqual(list(line.plots[0].series[0].values), LINE["series"][0]["values"])
        self.assertEqual(list(line.plots[0].series[1].values), LINE["series"][1]["values"])   # пропуск остаётся пропуском
        bar = found[1]
        self.assertEqual(list(bar.plots[0].series[0].values), BAR["series"][0]["values"])
        waterfall = found[2]
        base, shown = waterfall.plots[0].series
        self.assertEqual(list(waterfall.plots[0].categories), ["Начало", "Чеки", "Средний чек", "Итог"])
        self.assertEqual(list(shown.values), [12030.0, 420.5, 130.2, 12320.3])
        self.assertEqual(list(base.values), [0.0, 12030.0, 12320.3, 0.0])
        scatter = found[3]
        self.assertEqual(list(scatter.plots[0].series[0].values), [18.2, 27.0])

    def test_charts_are_editable(self):
        """Нативные диаграммы: у каждой — встроенная книга Excel («Изменить данные» в PowerPoint)."""
        blob, prs = deck(AGENT)
        with zipfile.ZipFile(io.BytesIO(blob)) as archive:
            embedded = [n for n in archive.namelist() if n.startswith("ppt/embeddings/") and n.endswith(".xlsx")]
        self.assertEqual(len(embedded), len(charts(prs)))

    def test_kpi_and_totals_are_code_numbers(self):
        _blob, prs = deck(AGENT)
        head = texts(prs.slides[1])
        self.assertIn("10 812 400", head)
        self.assertIn("−437 600 (−3,9 %)", head)
        compare = texts(prs.slides[2])
        self.assertIn("1 326 182,5 ₽", compare)
        self.assertIn("+57,0 %", compare)

    def test_tables_are_pptx_tables_with_russian_numbers(self):
        _blob, prs = deck(AGENT)
        tables = [shape.table for slide in prs.slides for shape in slide.shapes if shape.has_table]
        main = tables[0]
        self.assertEqual([c.text for c in main.rows[0].cells], AGENT["columns"])
        self.assertEqual([c.text for c in main.rows[1].cells], ["01.09.2026", "433 380,5"])
        self.assertEqual([c.text for c in main.rows[2].cells], ["02.09.2026", "451 402,0"])

    def test_waterfall_axis_note(self):
        _blob, prs = deck(AGENT)
        slide = next(s for s in prs.slides if "Факторы изменения" in texts(s))
        self.assertIn("Ось значений начинается с", texts(slide))
        crossing = dict(WATERFALL, start=100.0, series=[{"name": "Вклад", "values": [-300.0, 50.0]}], total=-150.0)
        self.assertIsNone(slides._waterfall_bars(crossing))       # через ноль — обычные столбики

    def test_sources_slide_and_no_sql(self):
        blob, prs = deck(AGENT)
        last = texts(prs.slides[-1])
        for expected in ("1–3 сентября 2026 и 2025", "все АЗС общества", "План сентябрь.xlsx", "АУП общества (Юг)",
                         "ОНПО «Юг» — 312 АЗС (по ОХД)", "сообщение 11, запись журнала 501",
                         "25.09.2026 14:05 МСК", "dm.data_for_ai_analytic_part_1", "r1 — выручка по дням"):
            self.assertIn(expected, last)
        with zipfile.ZipFile(io.BytesIO(blob)) as archive:
            xml = " ".join(archive.read(n).decode("utf-8", "ignore") for n in archive.namelist() if n.endswith(".xml"))
        self.assertNotIn("SELECT", xml)
        self.assertNotIn("sum_receipt_netto_ntu", xml)

    def test_recommendations_are_marked(self):
        _blob, prs = deck(AGENT)
        slide = next(s for s in prs.slides if "Что можно сделать" in texts(s))
        text = texts(slide)
        self.assertIn("РЕКОМЕНДАЦИЯ", text)
        self.assertIn("Основание:", text)
        self.assertIn("справочный характер", text)

    def test_potx_template(self):
        """Р-11: официальный шаблон .potx — его размеры и мастер; красную полосу проекта не рисуем."""
        with tempfile.TemporaryDirectory() as tmp:
            base = Presentation()
            base.slide_width, base.slide_height = Inches(10.833), Inches(7.5)
            buffer = io.BytesIO()
            base.save(buffer)
            potx = pathlib.Path(tmp) / "onpo.potx"
            with zipfile.ZipFile(io.BytesIO(buffer.getvalue())) as src, zipfile.ZipFile(potx, "w") as dst:
                for item in src.infolist():
                    data = src.read(item.filename)
                    if item.filename == "[Content_Types].xml":
                        data = data.replace(b"presentationml.presentation.main+xml", b"presentationml.template.main+xml")
                    dst.writestr(item, data)
            _blob, prs = deck(FAST, template=str(potx))
        self.assertEqual(prs.slide_width, Inches(10.833))
        self.assertEqual(len(prs.slides), 3)
        first = prs.slides[0]
        self.assertFalse(any(s.top == 0 and s.left == 0 and s.width == prs.slide_width for s in first.shapes))

    def test_number_format(self):
        self.assertEqual(slides.fmt_number(-437600.0), "−437 600")
        self.assertEqual(slides.fmt_number(3.9, signed=True), "+3,9")
        self.assertEqual(slides.fmt_number(0.0, 1), "0,0")
        self.assertEqual(slides.fmt_number(None), "—")


class _User:
    def __init__(self, uid, role):
        self.id, self.role, self.isAdmin = uid, role, False
        self.email = f"user{uid}@example.com"
        self.name, self.roleTitle, self.roleBinding, self.scopeLabel = "Тест", "", "", ""
        self.aiDialog = True


class EndpointTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self._saved = (journal.JOURNAL_DB, store.JOURNAL_DB, os.environ.get("AI_DEMO_ENABLED"))
        os.environ["AI_DEMO_ENABLED"] = "1"
        journal.JOURNAL_DB = store.JOURNAL_DB = pathlib.Path(self._tmp.name) / "journal.db"
        self.current = _User(7, "aup_npo")
        app = FastAPI()
        app.include_router(ai_api.build_router(lambda: self.current, lambda: self.current))
        self.client = TestClient(app)
        self.own = self._message(7, AGENT)
        self.other = self._message(8, FAST)
        self.draft = self._message(7, dict(FAST, planCard={"status": "draft"}))

    def tearDown(self):
        journal.JOURNAL_DB, store.JOURNAL_DB, flag = self._saved
        if flag is None:
            os.environ.pop("AI_DEMO_ENABLED", None)
        else:
            os.environ["AI_DEMO_ENABLED"] = flag
        self._tmp.cleanup()

    def _message(self, user_id, answer):
        journal_id = journal.write({"actor": f"user{user_id}@example.com", "role": "aup_npo", "binding": "Юг",
                                    "scope_label": answer["scopeLabel"], "question": answer["question"],
                                    "sql_final": SQL, "verdict": "ok", "row_count": len(answer["rows"])})
        dialog = store.create_dialog(user_id)["id"]
        return store.append_message(dialog, user_id, answer["question"], answer, journal_id)["id"]

    def test_own_answer_downloads_and_is_logged(self):
        response = self.client.get(f"/api/ai/messages/{self.own}/presentation")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["content-type"], slides.MEDIA_TYPE)
        self.assertIn(".pptx", response.headers["content-disposition"])
        prs = Presentation(io.BytesIO(response.content))
        last = texts(prs.slides[-1])
        self.assertIn("dm.data_for_ai_analytic_part_1", last)       # таблицы — из SQL журнала, сам SQL — нет
        self.assertIn("(Юг)", last)
        self.assertNotIn("SELECT", last)
        conn = store._connect()
        try:
            row = conn.execute("SELECT part, format, row_count FROM ai_exports WHERE message_id = ?", (self.own,)).fetchone()
        finally:
            conn.close()
        self.assertEqual((row["part"], row["format"], row["row_count"]), ("deck", "pptx", len(prs.slides)))

    def test_foreign_answer_is_not_found(self):
        self.assertEqual(self.client.get(f"/api/ai/messages/{self.other}/presentation").status_code, 404)

    def test_plan_draft_is_refused(self):
        self.assertEqual(self.client.get(f"/api/ai/messages/{self.draft}/presentation").status_code, 409)


if __name__ == "__main__":
    unittest.main()
