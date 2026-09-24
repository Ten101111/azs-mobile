"""ИИ-12: выгрузка таблиц и данных графиков ответа в XLSX и CSV с паспортом.

Модели и витрины здесь нет: ответ кладётся в хранилище диалогов так же, как его
кладёт /ask (response.model_dump()), вместе с записью журнала.
"""
import io
import os
import pathlib
import tempfile
import unittest
from datetime import date

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.ai import api as ai_api
from backend.ai import dialogs as store
from backend.ai import export as ai_export
from backend.ai import journal, pipeline

SQL = ("SELECT f.account_date AS \"Дата\", f.num_azs AS \"АЗС\", SUM(f.sum_receipt_netto_ntu) AS \"Выручка НТУ, ₽\" "
       "FROM dm.data_for_ai_analytic_part_1 AS f WHERE f.account_date >= DATE '2026-09-01' "
       "AND f.account_date <= DATE '2026-09-20' GROUP BY 1, 2")
FAST = {"ok": True, "question": "Выручка НТУ по дням", "scopeLabel": "ОНПО «Юг» — 12 АЗС (по ОХД)",
        "columns": ["Дата", "АЗС", "Выручка НТУ, ₽", "Конверсия НТУ, %"],
        "rows": [["2026-09-01", "10005", 125000.5, 42.93], ["2026-09-02", "10005", 118300, 41.1]],
        "truncated": True, "depth": "fast", "taskType": "lookup", "frame": {}, "steps": [], "tables": [], "charts": []}
AGENT = {"ok": True, "question": "Как менялась конверсия", "scopeLabel": "вся сеть", "depth": "analyze",
         "taskType": "trend", "columns": ["День", "Конверсия НТУ, %"], "rows": [["2026-09-01", 42.9]],
         "frame": {"mainResult": "r1", "period": "сентябрь 2026", "filters": ["только ОНПО «Юг»"],
                   "metrics": ["conversion_ntu"]},
         "steps": [{"key": "s1", "kind": "sql", "resultId": "r1", "purpose": "Конверсия по дням", "sql": SQL}],
         "tables": [{"id": "p1", "title": "Разложение по ОНПО", "columns": ["ОНПО", "Вклад, п. п."],
                     "rows": [["Юг", -0.4], ["Север", 0.1]], "truncated": False}],
         "charts": [{"id": "c1", "type": "line", "title": "Конверсия по дням", "source": "r1", "unit": "%",
                     "xTitle": "День", "x": ["01.09", "02.09"], "series": [{"name": "Конверсия НТУ", "values": [42.9, 41.1]}]},
                    {"id": "c2", "type": "kpi", "title": "Итоги", "source": "p1",
                     "cards": [{"label": "Конверсия", "value": 42.1, "delta": -0.8}]}]}


class _User:
    def __init__(self, uid, role, is_admin=False):
        self.id, self.role, self.isAdmin = uid, role, is_admin
        self.email = f"user{uid}@example.com"
        self.name, self.roleTitle, self.roleBinding, self.scopeLabel = "Тест", "", "", ""
        self.aiDialog = True


class ExportTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self._saved = (journal.JOURNAL_DB, store.JOURNAL_DB, os.environ.get("AI_DEMO_ENABLED"))
        os.environ["AI_DEMO_ENABLED"] = "1"
        path = pathlib.Path(self._tmp.name) / "journal.db"
        journal.JOURNAL_DB = store.JOURNAL_DB = path
        self.current = _User(7, "aup_npo")
        app = FastAPI()
        app.include_router(ai_api.build_router(lambda: self.current, lambda: self.current))
        self.client = TestClient(app)
        self.fast = self._message(7, FAST, role="aup_npo", binding="Юг")
        self.agent = self._message(7, AGENT, role="aup_npo", binding="Юг")

    def tearDown(self):
        journal.JOURNAL_DB, store.JOURNAL_DB, flag = self._saved
        if flag is None:
            os.environ.pop("AI_DEMO_ENABLED", None)
        else:
            os.environ["AI_DEMO_ENABLED"] = flag
        self._tmp.cleanup()

    def _message(self, user_id, answer, role, binding):
        journal_id = journal.write({"actor": f"user{user_id}@example.com", "role": role, "binding": binding,
                                    "scope_label": answer["scopeLabel"], "question": answer["question"],
                                    "sql_final": SQL, "verdict": "ok", "row_count": len(answer["rows"])})
        dialog = store.create_dialog(user_id)["id"]
        return store.append_message(dialog, user_id, answer["question"], answer, journal_id)["id"]

    def get(self, message_id, part="main", fmt="xlsx"):
        return self.client.get(f"/api/ai/messages/{message_id}/export", params={"part": part, "format": fmt})

    def test_xlsx_has_data_and_passport(self):
        response = self.get(self.fast)
        self.assertEqual(response.status_code, 200)
        self.assertIn("filename*=UTF-8''", response.headers["content-disposition"])
        book = load_workbook(io.BytesIO(response.content))
        self.assertEqual(book.sheetnames, ["Данные", "Паспорт"])
        data = book["Данные"]
        self.assertEqual([c.value for c in data[1]], FAST["columns"])
        self.assertEqual(data.cell(2, 1).value.date(), date(2026, 9, 1))      # дата — датой
        self.assertEqual(data.cell(2, 1).number_format, "DD.MM.YYYY")
        self.assertEqual(data.cell(2, 3).value, 125000.5)                    # число — числом
        self.assertEqual(data.cell(2, 3).number_format, "#,##0.0")
        self.assertEqual(data.cell(2, 4).number_format, "#,##0.00")
        self.assertEqual(data.cell(2, 2).value, "10005")                     # номер АЗС остаётся текстом
        self.assertIn("Конфиденциально", data.oddHeader.right.text)
        passport = {row[0].value: row[1].value for row in book["Паспорт"].iter_rows(min_row=2)}
        self.assertEqual(passport["Вопрос"], "Выручка НТУ по дням")
        self.assertIn("Конфиденциально", passport["Гриф"])
        self.assertIn("АУП общества (Юг)", passport["Роль и область данных"])
        self.assertIn("ОНПО «Юг»", passport["Роль и область данных"])
        self.assertIn("dm.data_for_ai_analytic_part_1", passport["Источник"])
        self.assertEqual(passport["Период"], "по условию запроса: 01.09.2026 — 20.09.2026")
        self.assertIn("выгрузка неполная", passport["Строк в выгрузке"])
        self.assertIn("факт", passport["Тип данных"])
        self.assertIn(f"сообщение {self.fast}", passport["Номер ответа"])
        self.assertNotIn("SQL (администратор и субадминистратор)", passport)  # АУП запрос не видит

    def test_sql_in_passport_only_for_admin(self):
        self.current = _User(9, "admin", is_admin=True)
        own = self._message(9, dict(FAST, sql=SQL), role="admin", binding="")
        book = load_workbook(io.BytesIO(self.get(own).content))
        passport = {row[0].value: row[1].value for row in book["Паспорт"].iter_rows(min_row=2)}
        self.assertEqual(passport["SQL (администратор и субадминистратор)"], SQL)

    def test_csv_for_russian_excel(self):
        response = self.get(self.fast, fmt="csv")
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b"\xef\xbb\xbf"))            # BOM
        text = response.content.decode("utf-8-sig")
        lines = text.split("\r\n")
        self.assertEqual(lines[0], "Конфиденциально — выгрузка ответа ИИ-аналитика")
        self.assertEqual(lines[2], "Дата;АЗС;Выручка НТУ, ₽;Конверсия НТУ, %")
        self.assertEqual(lines[3], "01.09.2026;10005;125000,5;42,93")
        self.assertIn("Паспорт выгрузки", text)
        self.assertIn("Период;по условию запроса: 01.09.2026 — 20.09.2026", text)

    def test_tables_and_charts_of_the_agent_answer(self):
        self.assertEqual(ai_export.parts(AGENT), ["main", "table-p1", "chart-c1", "chart-c2"])
        book = load_workbook(io.BytesIO(self.get(self.agent, "chart-c1").content))
        data = book["Данные"]
        self.assertEqual([c.value for c in data[1]], ["День", "Конверсия НТУ, %"])
        self.assertEqual([c.value for c in data[3]], ["02.09", 41.1])
        passport = {row[0].value: row[1].value for row in book["Паспорт"].iter_rows(min_row=2)}
        self.assertEqual(passport["Период"], "сентябрь 2026")
        self.assertIn("только ОНПО «Юг»", passport["Фильтры и параметры"])
        self.assertIn("данные графика: «Конверсия по дням»", passport["Что выгружено"])
        calc = load_workbook(io.BytesIO(self.get(self.agent, "table-p1").content))
        passport = {row[0].value: row[1].value for row in calc["Паспорт"].iter_rows(min_row=2)}
        self.assertIn("расчёт", passport["Тип данных"])
        kpi = load_workbook(io.BytesIO(self.get(self.agent, "chart-c2").content))["Данные"]
        self.assertEqual([c.value for c in kpi[1]], ["Показатель", "Значение", "Изменение"])

    def test_foreign_answers_unknown_parts_and_formats_are_refused(self):
        stranger = self._message(8, FAST, role="aup_npo", binding="Север")
        self.assertEqual(self.get(stranger).status_code, 404)                # чужой ответ
        self.assertEqual(self.get(self.agent, "table-r9").status_code, 404)
        self.assertEqual(self.get(self.fast, fmt="pdf").status_code, 400)

    def test_every_export_is_recorded(self):
        self.get(self.fast)
        self.get(self.agent, "chart-c1", "csv")
        conn = store._connect()
        try:
            rows = [tuple(r) for r in conn.execute("SELECT user_id, message_id, part, format, row_count FROM ai_exports")]
        finally:
            conn.close()
        self.assertEqual(rows, [(7, self.fast, "main", "xlsx", 2), (7, self.agent, "chart-c1", "csv", 2)])

    def test_sql_does_not_leak_through_the_frame(self):
        answer = pipeline.Answer(ok=True, question="q", scope_label="s", sql="SELECT 1",
                                 frame={"sql": "SELECT 1", "period": "сентябрь"})
        self.assertNotIn("sql", ai_api._response(answer, show_sql=False).frame)
        self.assertEqual(ai_api._response(answer, show_sql=True).frame["sql"], "SELECT 1")


if __name__ == "__main__":
    unittest.main()
