"""Справки (СП-01…СП-05, СП-07): периоды, расчёт по витрине ОХД, план месяца, полнота, версии, файлы, доступ.

Данные — синтетическая витрина ОХД (backend/tests/report_stand.py): те же таблицы
и столбцы, что dm.data_for_ai_analytic_part_1/2, каталог в диалекте PostgreSQL,
валидатор с проекцией таблиц; проверенный SQL переводится в SQLite и выполняется
на временном файле. Особые объекты: 7005 — падение топлива на 40 %, 7007 — три
дня без продаж, 7009 — нет двух дней отчётной недели, 7012 — открыт в 2026 году.
"""
from __future__ import annotations

import json
import pathlib
import tempfile
import unittest
from datetime import date

from fastapi import FastAPI, HTTPException, Request
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend import summary
from backend.ai.catalog import Catalog
from backend.reports import fmt, periods, registry, service, store, weekly
from backend.tests import report_stand
from backend.tests.report_stand import ONPO, REPORT_WEEK


class StandMixin:
    stand_kwargs: dict = {}

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.folder = pathlib.Path(self._tmp.name)
        self.mart = report_stand.build(self.folder, **self.stand_kwargs)
        self.catalog = report_stand.mart_catalog()
        self.run = report_stand.runner(self.mart)
        report_stand.patch_validator(self, self.catalog)

    def tearDown(self):
        self._tmp.cleanup()

    def build(self, week=REPORT_WEEK, **kw):
        return weekly.build(week, catalog=self.catalog, run=self.run, **kw)


class PeriodTests(unittest.TestCase):
    def test_week_bounds_labels_and_year_shift(self):
        week = periods.last_complete_week(date(2026, 9, 21))  # понедельник
        self.assertEqual((week.start, week.end), (date(2026, 9, 14), date(2026, 9, 20)))
        self.assertEqual(periods.last_complete_week(date(2026, 9, 20)).end, date(2026, 9, 13))  # воскресенье — неделя ещё идёт
        self.assertEqual(week.iso, "2026-W38")
        self.assertEqual(week.label, "14–20 сентября 2026")
        ly = periods.last_year(week)
        self.assertEqual(ly.start.weekday(), 0)
        self.assertEqual(ly.start, date(2025, 9, 15))
        self.assertEqual(periods.Week(date(2026, 9, 28), date(2026, 10, 4)).label, "28 сентября – 4 октября 2026")
        self.assertEqual(periods.Week(date(2025, 12, 29), date(2026, 1, 4)).label, "29 декабря 2025 – 4 января 2026")
        self.assertEqual(periods.parse_iso("2026-W01").start, date(2025, 12, 29))
        self.assertEqual([h["date"] for h in periods.Week(date(2025, 12, 29), date(2026, 1, 4)).holidays()],
                         ["2026-01-01", "2026-01-02", "2026-01-03", "2026-01-04"])
        self.assertEqual(len(periods.trailing(week, 8)), 8)

    def test_russian_number_format(self):
        self.assertEqual(fmt.number(1234567.891, 1), "1 234 567,9")
        self.assertEqual(fmt.delta(-3.14), "−3,1 %")
        self.assertEqual(fmt.delta(0.4, share=True), "+0,4 п. п.")


class BuildTests(StandMixin, unittest.TestCase):
    def test_sections_and_totals(self):
        model = self.build()
        codes = [m["code"] for m in model["metrics"]]
        self.assertEqual(codes, ["fuel_volume", "revenue_total", "ntu_revenue", "ntu_vd", "checks_total",
                                 "fuel_checks", "ntu_checks", "ntu_avg_check", "conversion"])
        self.assertEqual(model["fuelUnit"], "т")
        self.assertIn("dm.data_for_ai_analytic_part_1", model["passport"]["source"])
        self.assertEqual({r["name"] for r in model["onpo"]}, set(ONPO))
        total = next(m for m in model["metrics"] if m["code"] == "fuel_volume")["value"]
        self.assertAlmostEqual(sum(r["fuel"] for r in model["onpo"]), total, delta=len(model["onpo"]))
        self.assertEqual(model["passport"]["stationsWeek"], 36)
        self.assertEqual(model["week"]["iso"], "2026-W38")
        self.assertEqual(len(model["dynamics"]["charts"]), 2)
        self.assertTrue(all(len(s["values"]) == 8 for c in model["dynamics"]["charts"] for s in c["series"]))
        self.assertTrue(model["headline"][0].startswith("Реализация топлива за неделю"))
        self.assertEqual(model["pending"], [])  # план и сервис в выпуске — ничего не ждём

    def test_comparable_base_and_attention(self):
        model = self.build()
        p = model["passport"]
        self.assertEqual((p["comparablePrev"], p["excludedPrev"]), (35, 1))   # S09 без двух дней
        self.assertEqual((p["comparableYear"], p["excludedYear"]), (34, 2))   # S09 и S12 без прошлого года
        drops = [r["label"] for r in model["attention"]["drops"]]
        self.assertIn("АЗС № 10005", drops)
        idle = {r["label"]: r["idleDays"] for r in model["attention"]["noSales"]}
        self.assertEqual(idle.get("АЗС № 10007"), 3)
        # Падение S05 не тянет вниз сопоставимую базу больше, чем его доля: изменение считается по 35 объектам.
        fuel = next(m for m in model["metrics"] if m["code"] == "fuel_volume")
        self.assertLess(fuel["deltaPrev"], 0)
        conversion = next(m for m in model["metrics"] if m["code"] == "conversion")
        self.assertTrue(conversion["isShare"])

    def test_week_equals_summary_for_the_same_dates(self):
        model = self.build()
        from unittest import mock
        from backend.ai.validator import Scope, validate

        tiles = [t for t in summary.TILES if t.code in ("fuel_volume", "ntu_revenue", "checks_total")]
        with mock.patch.object(summary, "CATALOG", self.catalog):
            sql = summary.build_query(tiles, REPORT_WEEK.start.isoformat(), REPORT_WEEK.end.isoformat())
        result = self.run(validate(sql, Scope.all_network()).sql, 10)
        svod = dict(zip(result.columns, result.rows[0]))
        by_code = {m["code"]: m["value"] for m in model["metrics"]}
        for code in ("fuel_volume", "ntu_revenue", "checks_total"):
            self.assertAlmostEqual(by_code[code], float(svod[code]), places=0)

    def test_missing_sunday_blocks_publication(self):
        with self.assertRaises(weekly.ReportError) as caught:
            self.build(week=REPORT_WEEK.shift(7))
        self.assertIn("Нет данных за 27.09.2026", str(caught.exception))


class PlanTests(StandMixin, unittest.TestCase):
    """План месяца из part_2: % плана, % дней, отставание, темп; переходная неделя; нет плана — нет блока."""

    def test_open_month_plan_pace_and_status(self):
        model = self.build()
        month, = model["plan"]["months"]
        self.assertEqual((month["label"], month["closed"], month["daysPassed"], month["daysTotal"]),
                         ("сентябрь 2026", False, 20, 30))
        self.assertEqual(month["daysPct"], 66.7)
        rows = {r["code"]: r for r in month["rows"]}
        self.assertEqual(set(rows), {"fuel", "ntu", "vd"})
        ntu = rows["ntu"]
        self.assertAlmostEqual(ntu["pctMonth"], round(100 * ntu["fact"] / ntu["planMonth"], 1), places=1)
        self.assertEqual(ntu["gapPp"], round(ntu["pctMonth"] - 66.7, 1))
        self.assertAlmostEqual(ntu["needPerDay"], (ntu["planMonth"] - ntu["fact"]) / 10, delta=1)
        self.assertEqual(ntu["status"], weekly.plan_status(ntu["gapPp"]))
        self.assertEqual({o["name"] for o in month["onpo"]}, set(ONPO))
        self.assertTrue(any(line.startswith("План сентября 2026 по выручке НТУ") for line in model["headline"]))

    def test_transition_week_closes_the_old_month(self):
        model = self.build(week=periods.Week(date(2026, 8, 31), date(2026, 9, 6)))
        old, new = model["plan"]["months"]
        self.assertTrue(old["closed"])
        self.assertIsNone(next(r for r in old["rows"] if r["code"] == "ntu")["needPerDay"])
        self.assertEqual((new["label"], new["daysPassed"]), ("сентябрь 2026", 6))
        self.assertTrue(any(line.startswith("Итоги августа 2026") for line in model["headline"]))

    def test_status_thresholds(self):
        self.assertEqual([weekly.plan_status(g) for g in (1.0, -3.0, -3.1, -8.0, -8.1, None)],
                         ["норма", "норма", "внимание", "внимание", "критично", None])


class NoPlanTests(StandMixin, unittest.TestCase):
    stand_kwargs = {"plan_until": None}

    def test_without_plans_the_block_is_absent_and_the_reason_is_in_the_passport(self):
        model = self.build()
        self.assertIsNone(model["plan"])
        self.assertIn("План месяца и требуемый темп", model["pending"])
        self.assertIn("нет плана на сентябрь 2026", model["passport"]["planNote"])


class PartialPlanTests(StandMixin, unittest.TestCase):
    stand_kwargs = {"plan_until": date(2026, 9, 20)}

    def test_plan_loaded_only_to_date_gives_plan_to_date_without_pace(self):
        model = self.build()
        ntu = next(r for r in model["plan"]["months"][0]["rows"] if r["code"] == "ntu")
        self.assertIsNone(ntu["pctMonth"])
        self.assertIsNone(ntu["needPerDay"])
        self.assertIsNotNone(ntu["pctToDate"])
        self.assertIn("не на все дни", model["plan"]["note"])


class ServiceTests(StandMixin, unittest.TestCase):
    """Сервис (СП-07): оценки в приложении, негатив, жалобы ЕГЛ — неделя, месяц, ОНПО, зона внимания."""

    def setUp(self):
        super().setUp()
        self.model = self.build()
        self.service = self.model["service"]
        self.rows = {r["code"]: r for r in self.service["metrics"]}

    def test_week_values_and_changes_on_the_comparable_base(self):
        ratings = sum(7 * (10 + i % 5) for i in range(1, 37))
        points = sum(7 * (5 * (10 + i % 5) - 4) for i in range(1, 37)) - 4 * 4 - 3 * 3  # «1» ×4, «2» ×3
        self.assertEqual(self.rows["ratings"]["value"], ratings)
        self.assertEqual(self.rows["avg"]["value"], round(points / ratings, 3))
        self.assertEqual((self.rows["negative"]["value"], self.rows["negative"]["prev"],
                          self.rows["negative"]["lastYear"]), (7, 1, 1))
        self.assertEqual(self.rows["negative"]["deltaPrev"], 6)   # разность штук, а не +600 %
        self.assertEqual(self.rows["negative"]["deltaKind"], "abs")
        self.assertEqual(self.rows["complaints"]["value"], 1)
        self.assertEqual(self.rows["negativeShare"]["value"], round(100 * 7 / ratings, 2))
        checks = next(m["value"] for m in self.model["metrics"] if m["code"] == "checks_total")
        self.assertAlmostEqual(self.rows["quality"]["value"], round(100000 * 8 / checks, 2), places=2)
        self.assertLess(self.rows["avg"]["deltaPrev"], 0)

    def test_month_to_date_categories_and_onpo(self):
        month = self.service["months"][0]
        self.assertEqual((month["label"], month["from"], month["to"], month["closed"]),
                         ("сентябрь 2026", "2026-09-01", "2026-09-20", False))
        self.assertEqual(month["negative"], 8)  # 7 за отчётную неделю и 1 за прошлую
        self.assertEqual([(c["title"], c["week"], c["prev"]) for c in self.service["categories"]][:2],
                         [("Обслуживание на кассе", 3, 1), ("Техническое состояние", 2, 0)])
        avgs = [o["avg"] for o in self.service["onpo"]]
        self.assertEqual(avgs, sorted(avgs))  # слабые — сверху
        self.assertEqual(sum(o["negative"] for o in self.service["onpo"]), 7)

    def test_negative_joins_the_attention_zone_and_headline(self):
        att = self.model["attention"]
        self.assertEqual([(r["label"], r["negative"], r["category"]) for r in att["negative"]],
                         [("АЗС № 10005", 4, "Обслуживание на кассе"), ("АЗС № 10020", 2, "Техническое состояние")])
        self.assertEqual(att["negativeTotal"], 2)  # у 7030 одна «1» — ниже порога
        text = " ".join(self.model["headline"])
        self.assertIn("В зоне внимания 3 объекта", text)  # 7005 — и падение, и негатив: считается один раз
        self.assertIn("2 и больше негативных оценок в приложении — 2", text)
        self.assertIn("Средняя оценка в приложении за неделю — ", text)
        self.assertIn("негативных оценок — 7, жалоб ЕГЛ — 1", text)
        self.assertNotIn("Сервис: оценки в приложении и жалобы", self.model["pending"])
        self.assertIn("не официальный уровень сервиса", " ".join(self.model["passport"]["rules"]))

    def test_source_names_the_service_mart(self):
        self.assertTrue(self.model["passport"]["source"].endswith("(планы и оценки сервиса)"))


class NoServiceTests(StandMixin, unittest.TestCase):
    stand_kwargs = {"ratings": False}

    def test_no_ratings_no_block_and_the_reason_is_in_the_passport(self):
        model = self.build()
        self.assertIsNone(model["service"])
        self.assertIn("Сервис: оценки в приложении и жалобы", model["pending"])
        self.assertIn("нет оценок клиентов за неделю", model["passport"]["serviceNote"])
        self.assertNotIn("negative", model["attention"])
        self.assertIsNotNone(model["plan"])  # план по-прежнему есть


class SparseTests(StandMixin, unittest.TestCase):
    stand_kwargs = {"sparse": True}

    def test_incomplete_week_is_not_published(self):
        with self.assertRaises(weekly.ReportError) as caught:
            self.build()
        self.assertIn("Данные неполные", str(caught.exception))
        model = self.build(require_complete=False)
        self.assertLess(model["passport"]["completeness"]["sharePct"], 95)


class ServerMixin(StandMixin):
    """Сервер справок: принимает готовые модели, рисует файлы, хранит версии."""

    def setUp(self):
        super().setUp()
        self.db, self.files = self.folder / "reports.sqlite3", self.folder / "files"

    def model(self, week=REPORT_WEEK):
        from backend.reports import narrative

        model = self.build(week=week)
        model["narrative"] = narrative.template(model)
        return model

    def accept(self, model=None, **kw):
        return service.accept(model or self.model(), files_dir=self.files, db_path=self.db, **kw)


class ServerTests(ServerMixin, unittest.TestCase):
    def test_accept_is_idempotent_and_reissue_needs_a_request(self):
        first = self.accept()
        self.assertEqual((first["outcome"], first["issue"]["version"]), ("published", 1))
        again = self.accept()
        self.assertEqual((again["outcome"], again["same"]), ("exists", True))
        request = service.request_reissue("2026-W38", "правка данных", created_by="a@example.com", db_path=self.db)
        self.assertEqual(service.status("2026-W38", db_path=self.db)["request"]["id"], request["request"]["id"])
        second = self.accept(request_id=request["request"]["id"])
        self.assertEqual(second["issue"]["version"], 2)
        with store.connect(self.db) as conn:
            statuses = [(r["version"], r["status"]) for r in store.archive(conn, "network_weekly", "network")]
            current = store.full(store.latest(conn, "network_weekly", "network"))
            self.assertIsNone(store.open_request(conn, "network_weekly", "network"))
        self.assertEqual(statuses, [(2, "published"), (1, "replaced")])
        self.assertEqual(current["model"]["reason"], "правка данных")
        self.assertTrue(any("Версия формул" in line for line in current["model"]["passportLines"]))

    def test_only_models_built_on_the_dwh_mart_are_accepted(self):
        model = self.model()
        model["passport"]["source"] = "База KPI платформы (из ОХД)"
        with self.assertRaises(service.InvalidModel):
            self.accept(model)
        broken = self.model()
        broken["week"]["from"] = "2026-09-13"
        with self.assertRaises(service.InvalidModel):
            self.accept(broken)

    def test_issue_not_from_the_mart_is_replaced_without_request(self):
        legacy = self.model()
        legacy["passport"]["source"] = "Витрина ОХД"  # так подписывала выпуск сборка по базе KPI до 24.09.2026
        with store.connect(self.db) as conn:
            store.publish(conn, legacy, None, None, version=1)
        service.request_reissue("2026-W38", "Выгрузка из ДВХ", db_path=self.db)
        out = self.accept()
        self.assertEqual((out["outcome"], out["issue"]["version"]), ("published", 2))
        self.assertIsNone(service.status("2026-W38", db_path=self.db)["request"])  # запрос выполнен этим выпуском
        with store.connect(self.db) as conn:
            current = store.full(store.latest(conn, "network_weekly", "network"))
        self.assertEqual(current["reason"], service.LEGACY_REASON)
        self.assertEqual(self.accept()["outcome"], "exists")  # выпуск по витрине уже не заменяется без запроса

    def test_request_made_after_the_build_stays_open(self):
        model = self.model()
        model["passport"]["generatedAt"] = "2026-09-24T08:00:00+03:00"
        service.request_reissue("2026-W38", "правка", db_path=self.db)
        self.assertEqual(self.accept(model)["outcome"], "published")
        self.assertIsNotNone(service.status("2026-W38", db_path=self.db)["request"])

    def test_not_formed_is_shown_until_the_week_arrives(self):
        self.accept()
        out = service.not_formed("2026-W39", "Нет данных за 27.09.2026", db_path=self.db)
        self.assertTrue(out["changed"])
        self.assertFalse(service.not_formed("2026-W39", "Нет данных за 27.09.2026", db_path=self.db)["changed"])
        with store.connect(self.db) as conn:
            self.assertEqual(store.pending(conn, "network_weekly", "network")["week"], "2026-W39")

    def test_files_match_the_model(self):
        issue = self.accept()["issue"]
        with store.connect(self.db) as conn:
            row = store.get(conn, issue["id"])
            model = json.loads(row["model_json"])
        pdf, xlsx = self.files / row["pdf_path"], self.files / row["xlsx_path"]
        self.assertTrue(pdf.read_bytes().startswith(b"%PDF"))
        self.assertLess(pdf.stat().st_size, 1_000_000)
        wb = load_workbook(xlsx)
        self.assertEqual(wb.sheetnames, ["Цифры недели", "План месяца", "Сервис", "ОНПО", "Зоны внимания",
                                         "Динамика 8 недель", "Паспорт"])
        service_sheet = wb["Сервис"]
        self.assertEqual(service_sheet.cell(2, 1).value, "Средняя оценка в приложении")
        self.assertEqual(service_sheet.cell(2, 2).value, model["service"]["metrics"][0]["value"])
        zones = [row[0].value for row in wb["Зоны внимания"].iter_rows(min_row=2)]
        self.assertEqual(zones.count("Негатив в приложении"), 2)
        sheet = wb["Цифры недели"]
        for index, metric in enumerate(model["metrics"], start=2):
            self.assertEqual(sheet.cell(index, 1).value, metric["title"])
            self.assertEqual(sheet.cell(index, 3).value, metric["value"])
            self.assertEqual(sheet.cell(index, 5).value, metric["deltaPrev"])
        self.assertEqual(service.file_name(model, 1, "pdf"), "Справка_сеть_2026-W38_v1.pdf")


class ApiTests(ServerMixin, unittest.TestCase):
    """Доступ, приём выпусков с токеном, файлы с русским именем, запрос на перевыпуск, письмо о несформированной."""

    def setUp(self):
        super().setUp()
        from backend.reports import api

        self._saved_store = (store.DB_PATH, store.FILES_DIR, service.today_msk)
        store.DB_PATH, store.FILES_DIR = self.db, self.files
        service.today_msk = lambda: date(2026, 9, 23)
        self.sent = []

        class User:
            def __init__(self, admin):
                self.isAdmin, self.email, self.role = admin, "a@example.com", "admin" if admin else "territory_manager"

        def require_user(request: Request):
            return User(request.headers.get("x-admin") == "1")

        def require_admin(request: Request):
            if request.headers.get("x-admin") != "1":
                raise HTTPException(status_code=403, detail="нет доступа")
            return User(True)

        def require_import(request: Request):
            if request.headers.get("authorization") != "Bearer t":
                raise HTTPException(status_code=401, detail="нужен токен")

        app = FastAPI()
        app.include_router(api.build_router(require_user, require_admin, require_import,
                                            lambda subject, text: self.sent.append(subject)))
        self.client = TestClient(app)
        self.token = {"authorization": "Bearer t"}

    def tearDown(self):
        store.DB_PATH, store.FILES_DIR, service.today_msk = self._saved_store
        super().tearDown()

    def post_model(self, **extra):
        return self.client.post("/api/internal/reports/import", headers=self.token, json={"model": self.model(), **extra})

    def test_non_admin_is_refused(self):
        self.assertEqual(self.client.get("/api/reports").status_code, 403)
        self.assertEqual(self.client.post("/api/reports/generate", json={"reason": "проверка"}).status_code, 403)

    def test_import_publishes_and_files_download_with_russian_names(self):
        self.assertEqual(self.client.post("/api/internal/reports/import", json={"model": {}}).status_code, 401)
        status = self.client.post("/api/internal/reports/status", headers=self.token, json={}).json()
        self.assertEqual((status["week"]["iso"], status["published"]), ("2026-W38", False))
        built = self.post_model().json()
        self.assertEqual(built["outcome"], "published")
        current = self.client.get("/api/reports", headers={"x-admin": "1"}).json()["current"]
        self.assertEqual(current["week"], "2026-W38")
        self.assertEqual(current["model"]["metrics"][0]["code"], "fuel_volume")
        pdf = self.client.get(f"/api/reports/{current['id']}/file/pdf", headers={"x-admin": "1"})
        self.assertEqual(pdf.status_code, 200)
        self.assertTrue(pdf.content.startswith(b"%PDF"))
        self.assertIn("filename*=utf-8''%D0%A1%D0%BF%D1%80%D0%B0%D0%B2%D0%BA%D0%B0", pdf.headers["content-disposition"])
        self.assertEqual(self.post_model().json()["outcome"], "exists")
        bad = self.client.post("/api/internal/reports/import", headers=self.token, json={"model": {"type": "x"}})
        self.assertEqual(bad.status_code, 422)

    def test_reissue_is_a_request_for_the_mac(self):
        self.post_model()
        self.assertEqual(self.client.post("/api/reports/generate", json={"reason": ""}, headers={"x-admin": "1"}).status_code, 400)
        out = self.client.post("/api/reports/generate", json={"reason": "правка данных за воскресенье", "week": "2026-W38"},
                               headers={"x-admin": "1"}).json()
        self.assertEqual(out["outcome"], "requested")
        overview = self.client.get("/api/reports", headers={"x-admin": "1"}).json()
        self.assertEqual(overview["request"]["week"], "2026-W38")
        status = self.client.post("/api/internal/reports/status", headers=self.token, json={"week": "2026-W38"}).json()
        second = self.post_model(requestId=status["request"]["id"]).json()
        self.assertEqual(second["issue"]["version"], 2)
        archive = self.client.get("/api/reports", headers={"x-admin": "1"}).json()["archive"]
        self.assertEqual([(a["version"], a["status"]) for a in archive], [(2, "published"), (1, "replaced")])

    def test_missing_data_notifies_admin_once_after_the_last_morning_attempt(self):
        from backend.reports import api
        saved = api.final_attempt
        api.final_attempt = lambda now: True
        try:
            body = {"week": "2026-W39", "reason": "Нет данных за 27.09.2026: последний день в витрине — 20.09.2026"}
            first = self.client.post("/api/internal/reports/not-formed", headers=self.token, json=body).json()
            self.client.post("/api/internal/reports/not-formed", headers=self.token, json=body)
        finally:
            api.final_attempt = saved
        self.assertEqual(first["outcome"], "not_formed")
        self.assertEqual(len(self.sent), 1)
        overview = self.client.get("/api/reports", headers={"x-admin": "1"}).json()
        self.assertIn("Нет данных за 27.09.2026", overview["pending"]["reason"])


class AccessTests(unittest.TestCase):
    def test_only_admin_sees_reports_for_now(self):
        class U:
            isAdmin = False
        self.assertFalse(registry.allowed(U()))
        U.isAdmin = True
        self.assertTrue(registry.allowed(U()))


class MartOnlyTests(unittest.TestCase):
    def test_report_refuses_any_source_but_the_dwh_mart(self):
        kpi = Catalog(tables={"station_kpi_daily": {"ksss", "metric_date", "fuel_volume"}},
                      facts_table="station_kpi_daily", date_column="metric_date")
        with self.assertRaises(weekly.ReportError) as caught:
            weekly.build(REPORT_WEEK, catalog=kpi, run=lambda sql, limit: None)
        self.assertIn("только по витрине ОХД", str(caught.exception))


if __name__ == "__main__":
    unittest.main()
