"""«Журнал ИИ» (25.09.2026): тематики вопросов, фильтры, карточка обращения, выгрузка.

Тематики: правила — код по семантическому слою и словарям; модель (здесь —
сценарий) разбирает то, что правила не узнали, и заводит новые темы. Журнал —
во временном файле, настоящий data/ai_journal.sqlite3 не трогается.
"""
from __future__ import annotations

import io
import json
import pathlib
import tempfile
import time
import unittest

from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient

from backend.ai import api as ai_api
from backend.ai import dialogs as store
from backend.ai import journal, journal_view, topics
from backend.ai.agent.llm import Reply

NOW = 1_790_000_000


class _Model:
    """Модель-сценарий: отвечает заданными темами, записывает, что ей показали."""

    def __init__(self, answers):
        self.answers = list(answers)
        self.seen: list[list[dict]] = []

    def chat(self, messages, tools=None, json_mode=False, max_tokens=None):
        self.seen.append(messages)
        answer = self.answers.pop(0)
        if isinstance(answer, Exception):
            raise answer
        return Reply(content=json.dumps({"topic": answer}, ensure_ascii=False))


class JournalCase(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self._saved = (journal.JOURNAL_DB, store.JOURNAL_DB)
        path = pathlib.Path(self._tmp.name) / "journal.sqlite3"
        journal.JOURNAL_DB = path
        store.JOURNAL_DB = path

    def tearDown(self):
        journal.JOURNAL_DB, store.JOURNAL_DB = self._saved
        self._tmp.cleanup()

    def add(self, question, *, role="admin", depth="fast", verdict="ok", task=None, trace=None, at=None,
            total_ms=20_000, actor="admin@example.com", rule=None):
        entry_id = journal.write({"actor": actor, "role": role, "question": question, "verdict": verdict,
                                  "depth": depth, "task_type": task, "rule": rule,
                                  "trace_json": json.dumps(trace, ensure_ascii=False) if trace else None})
        conn = journal._connect()
        conn.execute("UPDATE ai_queries SET created_at = ?, total_ms = ? WHERE id = ?",
                     (at or NOW - 3600, total_ms, entry_id))
        conn.commit()
        conn.close()
        return entry_id


class RulesTests(unittest.TestCase):
    def test_metric_family_slice_task_and_period(self):
        facets = topics.classify("Сравни август 2026 с августом 2025 по обществам: выручка НТУ и средний чек НТУ")
        self.assertEqual(facets["metric"], ["Выручка НТУ", "Средний чек НТУ"])
        self.assertEqual(facets["theme"], ["Выручка и продажи НТУ", "Конверсия и средний чек"])
        self.assertEqual(facets["slice"], ["Общества (ОНПО)"])
        self.assertEqual(facets["task"], ["Сравнение"])
        self.assertIn("Год к году", facets["period"])

    def test_longer_phrase_wins_and_generic_words_do_not_name_a_metric(self):
        self.assertEqual(topics.metrics_in_text("Какой средний чек НТУ?")[0][0], "avg_check_ntu")
        self.assertEqual(topics.metrics_in_text("Кто территориал по АЗС 77551?"), [])
        self.assertEqual(topics.classify("Кто территориал по АЗС 77551?")["theme"], ["Руководители и персонал"])

    def test_plan_of_agent_fills_gaps(self):
        trace = {"plan": {"standaloneQuestion": "Сравни выручку НТУ Черепановой и Григорьевского", "taskType": "compare",
                          "metrics": ["revenue_ntu"], "period": "текущий месяц"}}
        facets = topics.classify("А теперь их двоих?", None, json.dumps(trace, ensure_ascii=False))
        self.assertEqual(facets["theme"], ["Выручка и продажи НТУ"])
        self.assertEqual(facets["task"], ["Сравнение"])

    def test_unknown_metric_becomes_its_own_theme(self):
        self.assertEqual(topics.theme_of_metric("plan_fuel_b2c", "План топлива B2C"), "План и выполнение")
        self.assertEqual(topics.theme_of_metric("new_metric", "Новый показатель"), "Новый показатель")

    def test_question_without_theme_waits_for_model(self):
        self.assertEqual(topics.classify("какие у тебя будут рекомендации?")["theme"], [])

    def test_model_topic_is_cleaned(self):
        self.assertEqual(topics.clean_topic(" уточнение ответа за 2026 год. "), "Уточнение ответа за год")
        self.assertEqual(topics.clean_topic("x" * 60), "")


class SyncAndModelTests(JournalCase):
    def test_sync_is_idempotent_and_marks_unknown_for_model(self):
        known = self.add("Выполнение плана выручки НТУ по обществам в текущем месяце")
        unknown = self.add("Пересчитай ещё раз")
        self.assertEqual(topics.sync(now=NOW), 2)
        self.assertEqual(topics.sync(now=NOW), 0)
        self.assertEqual(topics.pending_count(), 1)
        conn = topics.connect()
        tags = topics.query_topics(conn, [known, unknown])
        conn.close()
        self.assertEqual(tags[known]["theme"], ["План и выполнение"])
        self.assertEqual(tags[unknown]["theme"], [topics.UNKNOWN_THEME])

    def test_model_reuses_existing_theme_and_creates_new_one(self):
        first = self.add("Пересчитай ещё раз")
        second = self.add("Какие будут рекомендации по итогам?")
        topics.sync(now=NOW)
        model = _Model(["Уточнение предыдущего ответа", "уточнение предыдущего ответа"])
        self.assertEqual(topics.label_pending(model=model, limit=5, now=NOW), 2)
        conn = topics.connect()
        tags = topics.query_topics(conn, [first, second])
        origin = conn.execute("SELECT origin FROM ai_topics WHERE facet = 'theme' AND value = ?",
                              ("Уточнение предыдущего ответа",)).fetchone()[0]
        conn.close()
        self.assertEqual(tags[first]["theme"], tags[second]["theme"])      # одна тема, без дубля по регистру
        self.assertEqual(origin, "model")
        self.assertIn("Уточнение предыдущего ответа", model.seen[1][1]["content"])   # новая тема видна модели
        self.assertEqual(topics.pending_count(), 0)

    def test_model_failure_retries_then_gives_up(self):
        self.add("Пересчитай ещё раз")
        topics.sync(now=NOW)
        for _ in range(topics.MODEL_TRIES):
            topics.label_pending(model=_Model([RuntimeError("Ollama недоступна")]), now=NOW)
        conn = topics.connect()
        state = conn.execute("SELECT model_state, model_tries FROM ai_topic_state").fetchone()
        conn.close()
        self.assertEqual(tuple(state), ("failed", topics.MODEL_TRIES))

    def test_theme_cap_sends_new_ideas_to_other(self):
        saved = topics.MAX_MODEL_THEMES
        topics.MAX_MODEL_THEMES = 1
        try:
            self.add("Вопрос один без темы")
            self.add("Вопрос два без темы")
            topics.sync(now=NOW)
            topics.label_pending(model=_Model(["Первая тема", "Вторая тема"]), limit=5, now=NOW)
            values = {item["value"] for item in topics.themes()}
            self.assertIn("Первая тема", values)
            self.assertNotIn("Вторая тема", values)
        finally:
            topics.MAX_MODEL_THEMES = saved

    def test_new_rules_revision_keeps_model_themes(self):
        entry = self.add("Пересчитай ещё раз")
        topics.sync(now=NOW)
        topics.label_pending(model=_Model(["Уточнение ответа"]), now=NOW)
        saved = topics.RULES_REV
        topics.RULES_REV = saved + 1
        try:
            self.assertEqual(topics.sync(now=NOW), 1)
        finally:
            topics.RULES_REV = saved
        conn = topics.connect()
        self.assertEqual(topics.query_topics(conn, [entry])[entry]["theme"], ["Уточнение ответа"])
        conn.close()


class RetentionTests(JournalCase):
    def test_topics_leave_with_the_journal_entry(self):
        from backend.ai import retention

        old = self.add("Выручка НТУ по сети", at=NOW - 400 * 86400)
        fresh = self.add("Топливо по сети", at=NOW - 3600)
        topics.sync(now=NOW)
        retention.cleanup(now=NOW)
        conn = topics.connect()
        left = {row[0] for row in conn.execute("SELECT DISTINCT query_id FROM ai_query_topics")}
        states = {row[0] for row in conn.execute("SELECT query_id FROM ai_topic_state")}
        conn.close()
        self.assertEqual(left, {fresh})
        self.assertNotIn(old, states)


class SearchTests(JournalCase):
    def setUp(self):
        super().setUp()
        self.fuel = self.add("Сравни объём топлива за сентябрь 2026 с сентябрём 2025", depth="analyze", task="compare")
        self.plan = self.add("Выполнение плана выручки НТУ по обществам в текущем месяце", role="aup_npo")
        self.error = self.add("Выручка НТУ по АЗС за август", verdict="execution_error")
        self.old = self.add("Выручка НТУ по сети", at=NOW - 40 * 86400)
        topics.sync(now=NOW)

    def test_filters_and_counts(self):
        data = journal_view.search(days=30, now=NOW)
        self.assertEqual(data["total"], 3)                        # старое — за пределами 30 дней
        self.assertEqual(data["summary"]["failed"], 1)
        theme = next(f for f in data["facets"] if f["code"] == "theme")
        counts = {v["value"]: v["count"] for v in theme["values"]}
        self.assertEqual(counts["Выручка и продажи НТУ"], 1)
        self.assertEqual(journal_view.search(days=30, outcome="failed", now=NOW)["entries"][0]["id"], self.error)
        self.assertEqual(journal_view.search(days=30, role="aup_npo", now=NOW)["entries"][0]["id"], self.plan)
        self.assertEqual(journal_view.search(days=30, depth="analyze", now=NOW)["entries"][0]["id"], self.fuel)

    def test_text_search_ignores_case_in_cyrillic(self):
        ids = [e["id"] for e in journal_view.search(days=0, text="ВЫРУЧКИ", now=NOW)["entries"]]
        self.assertEqual(ids, [self.plan])

    def test_topics_or_inside_and_between_facets(self):
        chosen = {"theme": ["Топливо", "План и выполнение"]}
        ids = {e["id"] for e in journal_view.search(days=30, chosen=chosen, now=NOW)["entries"]}
        self.assertEqual(ids, {self.fuel, self.plan})
        chosen["slice"] = ["Общества (ОНПО)"]
        data = journal_view.search(days=30, chosen=chosen, now=NOW)
        self.assertEqual([e["id"] for e in data["entries"]], [self.plan])
        # Счётчик своего фильтра считается без его выбора — видно, что даст соседнее значение.
        theme = {v["value"]: v["count"] for v in next(f for f in data["facets"] if f["code"] == "theme")["values"]}
        self.assertEqual(theme.get("Топливо", 0), 0)
        self.assertEqual(theme["План и выполнение"], 1)

    def test_new_values_are_marked(self):
        data = journal_view.search(days=0, now=NOW)
        theme = {v["value"]: v for v in next(f for f in data["facets"] if f["code"] == "theme")["values"]}
        self.assertTrue(theme["Топливо"]["isNew"])
        self.assertTrue(journal_view.search(days=0, now=NOW + 30 * 86400)["facets"][0]["values"][0]["isNew"] is False)

    def test_detail_has_steps_and_model_time(self):
        trace = {"plan": {"standaloneQuestion": "Выручка НТУ по обществам", "metrics": ["revenue_ntu"], "period": "август"},
                 "steps": [{"kind": "sql", "label": "Выручка по обществам", "ms": 1200, "ok": True, "sql": "SELECT 1"},
                           {"kind": "write", "label": "Сформулировал ответ", "ms": 9000, "ok": True}]}
        entry = self.add("Выручка НТУ по обществам", depth="analyze", trace=trace)
        journal.set_run_stats(entry, model_prompt_ms=3000, model_eval_ms=8000, model_load_ms=0,
                              model_trace=json.dumps({"calls": [{"in": 7000, "out": 300, "promptMs": 3000, "evalMs": 8000}]}))
        topics.sync(now=NOW)
        item = journal_view.detail(entry)
        self.assertEqual([s["kindTitle"] for s in item["steps"]], ["Запрос к витрине", "Итоговый текст"])
        self.assertNotIn("sql", item["steps"][0])                  # SQL шага не уходит в список шагов
        self.assertEqual((item["promptMs"], item["evalMs"]), (3000, 8000))
        self.assertEqual(item["planMetrics"], ["Выручка НТУ"])
        self.assertIsNone(journal_view.detail(999_999))


class ApiTests(JournalCase):
    def setUp(self):
        super().setUp()
        self.add("Сравни объём топлива за сентябрь 2026 с сентябрём 2025", at=int(time.time()) - 60)
        self.current = type("U", (), {"id": 1, "role": "admin", "isAdmin": True, "email": "a@b.c"})()

        def require_admin():
            if not getattr(self.current, "isAdmin", False):
                raise HTTPException(status_code=403, detail="Только администратор")
            return self.current

        app = FastAPI()
        app.include_router(ai_api.build_router(require_admin, lambda: self.current))
        self.client = TestClient(app)

    def test_list_detail_export_admin_only(self):
        data = self.client.get("/api/ai/admin/journal", params={"days": 7, "topic": ["theme:Топливо", "bad:x"]}).json()
        self.assertEqual(data["total"], 1)
        entry_id = data["entries"][0]["id"]
        self.assertEqual(self.client.get(f"/api/ai/admin/journal/{entry_id}").status_code, 200)
        self.assertEqual(self.client.get("/api/ai/admin/journal/999999").status_code, 404)
        blob = self.client.get("/api/ai/admin/journal/export", params={"days": 7})
        self.assertEqual(blob.status_code, 200)
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(blob.content)).active
        self.assertEqual(ws["A1"].value, "Дата и время")
        self.assertIn("топлива", ws["D2"].value)
        self.current.isAdmin = False
        self.assertEqual(self.client.get("/api/ai/admin/journal").status_code, 403)


if __name__ == "__main__":
    unittest.main()
