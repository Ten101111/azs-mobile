from __future__ import annotations

import hashlib
import json
import math
import os
import re
import tempfile
import zipfile
from datetime import date, datetime, time, timezone
from email.message import EmailMessage
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


PROJECT_DIR = Path(__file__).resolve().parents[1]
DATA_DIR = PROJECT_DIR / "data"
DEFAULT_XLSX_PATH = DATA_DIR / "fuel_outages_latest.xlsx"
DEFAULT_SNAPSHOT_PATH = DATA_DIR / "fuel_outages_latest.json"
DEFAULT_STALE_AFTER_SECONDS = 24 * 60 * 60
MAX_ARCHIVE_ENTRIES = 5000
MAX_UNCOMPRESSED_BYTES = 64 * 1024 * 1024
OUTAGE_ANALYTICS_GROUPS = {"region", "regionalManager", "territoryManager"}

DETAIL_COLUMNS = (
    "НПО",
    "Регион",
    "АЗС",
    "Код КССС",
    "Продукт",
    "Часы",
    "Дата",
    "Время начала",
    "Время окончания",
    "Ожид. реализ, л",
)

COLUMN_KEYS = {
    "НПО": "npo",
    "Регион": "region",
    "АЗС": "station",
    "Код КССС": "ksss",
    "Продукт": "product",
    "Часы": "hours",
    "Дата": "date",
    "Время начала": "startTime",
    "Время окончания": "endTime",
    "Ожид. реализ, л": "expectedSalesLiters",
}

NUMERIC_COLUMNS = {"Часы", "Ожид. реализ, л"}
TIME_COLUMNS = {"Время начала", "Время окончания"}


class FuelOutageImportError(Exception):
    def __init__(self, status_code: int, detail: str):
        super().__init__(detail)
        self.status_code = status_code
        self.detail = detail


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()


def _normalized_header(value: Any) -> str:
    return re.sub(r"[^0-9a-zа-яё]+", " ", _clean_text(value).casefold()).strip()


NORMALIZED_COLUMNS = {_normalized_header(column): column for column in DETAIL_COLUMNS}


def _number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        number = float(value)
        return number if math.isfinite(number) else None
    text = _clean_text(value).replace(" ", "").replace(",", ".")
    if not text:
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    return number if math.isfinite(number) else None


def _identifier(value: Any) -> str:
    text = _clean_text(value)
    if re.fullmatch(r"\d+[.,]0+", text):
        return text.split(".", 1)[0].split(",", 1)[0]
    return text


def _date_value(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _clean_text(value)
    for pattern in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(text, pattern).date().isoformat()
        except ValueError:
            continue
    return text


def _time_value(value: Any) -> str:
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        return value.isoformat(timespec="minutes")
    if isinstance(value, (int, float)) and 0 <= float(value) < 1:
        seconds = int(round(float(value) * 24 * 60 * 60)) % (24 * 60 * 60)
        return f"{seconds // 3600:02d}:{(seconds % 3600) // 60:02d}"
    return _clean_text(value)


def _cell_value(column: str, value: Any) -> Any:
    if column == "Код КССС":
        return _identifier(value)
    if column in NUMERIC_COLUMNS:
        return _number(value)
    if column == "Дата":
        return _date_value(value)
    if column in TIME_COLUMNS:
        return _time_value(value)
    return _clean_text(value)


def rows_from_matrix(matrix: Iterable[Iterable[Any]]) -> list[dict[str, Any]]:
    source_rows = [list(row) for row in matrix]
    header_index = -1
    column_indexes: dict[str, int] = {}

    for row_index, row in enumerate(source_rows[:50]):
        resolved = {
            NORMALIZED_COLUMNS[normalized]: index
            for index, value in enumerate(row)
            if (normalized := _normalized_header(value)) in NORMALIZED_COLUMNS
        }
        if len(resolved) == len(DETAIL_COLUMNS):
            header_index = row_index
            column_indexes = resolved
            break

    if header_index < 0:
        raise FuelOutageImportError(422, "Таблица 'Детально' не содержит ожидаемые колонки")

    result: list[dict[str, Any]] = []
    for row in source_rows[header_index + 1 :]:
        values = {
            column: _cell_value(column, row[index] if index < len(row) else None)
            for column, index in column_indexes.items()
        }
        if not any(value not in (None, "") for value in values.values()):
            continue
        if all(_normalized_header(values.get(column)) == _normalized_header(column) for column in DETAIL_COLUMNS):
            continue
        result.append({COLUMN_KEYS[column]: values[column] for column in DETAIL_COLUMNS})

    if not result:
        raise FuelOutageImportError(422, "Таблица 'Детально' не содержит строк данных")
    return result


def _validate_xlsx_archive(content: bytes) -> None:
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ARCHIVE_ENTRIES:
                raise FuelOutageImportError(413, "XLSX содержит слишком много файлов")
            if sum(entry.file_size for entry in entries) > MAX_UNCOMPRESSED_BYTES:
                raise FuelOutageImportError(413, "Распакованный XLSX превышает допустимый размер")
    except zipfile.BadZipFile as exc:
        raise FuelOutageImportError(422, "Передан поврежденный XLSX") from exc


def rows_from_xlsx(content: bytes) -> list[dict[str, Any]]:
    _validate_xlsx_archive(content)
    try:
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:
        raise FuelOutageImportError(422, "Не удалось открыть XLSX") from exc

    names = list(workbook.sheetnames)
    names.sort(key=lambda name: (name.casefold() != "детально", name.casefold()))
    errors: list[FuelOutageImportError] = []
    for name in names:
        worksheet = workbook[name]
        try:
            return rows_from_matrix(worksheet.iter_rows(values_only=True))
        except FuelOutageImportError as exc:
            errors.append(exc)
    raise errors[0] if errors else FuelOutageImportError(422, "В XLSX нет листов")


class _HtmlTableParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.tables: list[list[list[str]]] = []
        self._table: list[list[str]] | None = None
        self._row: list[str] | None = None
        self._cell: list[str] | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.casefold()
        if tag == "table" and self._table is None:
            self._table = []
        elif tag == "tr" and self._table is not None:
            self._row = []
        elif tag in {"td", "th"} and self._row is not None:
            self._cell = []
        elif tag == "br" and self._cell is not None:
            self._cell.append(" ")

    def handle_data(self, data: str) -> None:
        if self._cell is not None:
            self._cell.append(data)

    def handle_endtag(self, tag: str) -> None:
        tag = tag.casefold()
        if tag in {"td", "th"} and self._row is not None and self._cell is not None:
            self._row.append(_clean_text("".join(self._cell)))
            self._cell = None
        elif tag == "tr" and self._table is not None and self._row is not None:
            self._table.append(self._row)
            self._row = None
        elif tag == "table" and self._table is not None:
            self.tables.append(self._table)
            self._table = None


def rows_from_html(html_content: str) -> list[dict[str, Any]]:
    parser = _HtmlTableParser()
    parser.feed(html_content)
    errors: list[FuelOutageImportError] = []
    for table in parser.tables:
        try:
            return rows_from_matrix(table)
        except FuelOutageImportError as exc:
            errors.append(exc)
    raise errors[0] if errors else FuelOutageImportError(422, "В письме нет HTML-таблиц")


def rows_from_email(message: EmailMessage) -> list[dict[str, Any]]:
    attachment_errors: list[FuelOutageImportError] = []
    html_parts: list[str] = []
    for part in message.walk():
        filename = _clean_text(part.get_filename())
        content_type = part.get_content_type().casefold()
        if filename.casefold().endswith(".xlsx") or content_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
            payload = part.get_payload(decode=True) or b""
            if payload:
                try:
                    return rows_from_xlsx(payload)
                except FuelOutageImportError as exc:
                    attachment_errors.append(exc)
        elif content_type == "text/html":
            try:
                html_parts.append(part.get_content())
            except (LookupError, UnicodeDecodeError):
                payload = part.get_payload(decode=True) or b""
                html_parts.append(payload.decode("utf-8", errors="replace"))

    for html_content in html_parts:
        try:
            return rows_from_html(html_content)
        except FuelOutageImportError:
            continue
    if attachment_errors:
        raise attachment_errors[0]
    raise FuelOutageImportError(422, "В письме не найдена таблица 'Детально'")


def _safe_excel_text(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def xlsx_from_rows(rows: list[dict[str, Any]]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Детально"
    worksheet.freeze_panes = "A2"
    worksheet.append(list(DETAIL_COLUMNS))

    reverse_keys = {column: COLUMN_KEYS[column] for column in DETAIL_COLUMNS}
    for row in rows:
        worksheet.append([_safe_excel_text(row.get(reverse_keys[column])) for column in DETAIL_COLUMNS])

    header_fill = PatternFill("solid", fgColor="9B1C31")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    worksheet.row_dimensions[1].height = 24

    widths = (15, 24, 22, 14, 22, 12, 14, 16, 18, 18)
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[worksheet.cell(row=1, column=index).column_letter].width = width
    for row in worksheet.iter_rows(min_row=2):
        row[5].number_format = "0.00"
        row[9].number_format = '#,##0.00'

    if rows:
        table = Table(displayName="Detail", ref=f"A1:J{len(rows) + 1}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(table)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _atomic_write(path: Path, content: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(prefix=f".{path.name}.", dir=path.parent)
    temporary_path = Path(temporary_name)
    try:
        with os.fdopen(descriptor, "wb") as output:
            output.write(content)
            output.flush()
            os.fsync(output.fileno())
        temporary_path.chmod(0o600)
        os.replace(temporary_path, path)
        path.chmod(0o600)
    finally:
        temporary_path.unlink(missing_ok=True)


def outage_xlsx_path() -> Path:
    configured = os.getenv("FUEL_OUTAGE_XLSX_PATH", "").strip()
    return Path(configured) if configured else DEFAULT_XLSX_PATH


def outage_snapshot_path() -> Path:
    configured = os.getenv("FUEL_OUTAGE_SNAPSHOT_PATH", "").strip()
    return Path(configured) if configured else DEFAULT_SNAPSHOT_PATH


def stale_after_seconds() -> int:
    try:
        return max(60, int(os.getenv("FUEL_OUTAGE_STALE_AFTER_SECONDS", str(DEFAULT_STALE_AFTER_SECONDS))))
    except ValueError:
        return DEFAULT_STALE_AFTER_SECONDS


def _read_snapshot() -> dict[str, Any] | None:
    try:
        payload = json.loads(outage_snapshot_path().read_text(encoding="utf-8"))
    except (FileNotFoundError, OSError, json.JSONDecodeError):
        return None
    return payload if isinstance(payload, dict) else None


def _sum(values: Iterable[Any]) -> float:
    return round(sum(float(value) for value in values if isinstance(value, (int, float))), 2)


def is_outage_ongoing(item: dict[str, Any]) -> bool:
    end_time = _clean_text(item.get("endTime"))
    return not end_time or end_time[:5] == "23:59"


def aggregate_outage_snapshot(
    snapshot: dict[str, Any] | None,
    stations: Iterable[dict[str, Any]],
    group_by: str,
) -> dict[str, Any]:
    if group_by not in OUTAGE_ANALYTICS_GROUPS:
        raise ValueError("Unsupported fuel outage grouping")

    payload = snapshot or {}
    items = list(payload.get("items") or [])
    station_lookup = {
        _clean_text(station.get("ksss")): station
        for station in stations
        if _clean_text(station.get("ksss"))
    }
    group_fallbacks = {
        "region": "Регион не определен",
        "regionalManager": "РУ не определен",
        "territoryManager": "Территория не определена",
    }

    groups: dict[str, dict[str, Any]] = {}
    total_stations: set[str] = set()
    total_products: set[str] = set()
    unmatched_stations: set[str] = set()
    report_dates: list[str] = []

    for item in items:
        ksss = _clean_text(item.get("ksss"))
        station = station_lookup.get(ksss, {})
        if ksss:
            total_stations.add(ksss)
            if not station:
                unmatched_stations.add(ksss)

        if group_by == "region":
            label = _clean_text(item.get("region")) or _clean_text(station.get("subject"))
        else:
            label = _clean_text(station.get(group_by))
        label = label or group_fallbacks[group_by]

        group = groups.setdefault(
            label,
            {
                "id": label,
                "label": label,
                "eventCount": 0,
                "ongoingCount": 0,
                "totalHours": 0.0,
                "expectedSalesLiters": 0.0,
                "stations": set(),
                "products": set(),
            },
        )
        group["eventCount"] += 1
        group["ongoingCount"] += int(is_outage_ongoing(item))
        group["totalHours"] += float(item["hours"]) if isinstance(item.get("hours"), (int, float)) else 0.0
        group["expectedSalesLiters"] += (
            float(item["expectedSalesLiters"])
            if isinstance(item.get("expectedSalesLiters"), (int, float))
            else 0.0
        )
        if ksss:
            group["stations"].add(ksss)
        product = _clean_text(item.get("product"))
        if product:
            group["products"].add(product)
            total_products.add(product)
        report_date = _clean_text(item.get("date"))
        if report_date:
            report_dates.append(report_date)

    rows = []
    for group in groups.values():
        rows.append(
            {
                "id": group["id"],
                "label": group["label"],
                "eventCount": int(group["eventCount"]),
                "stationCount": len(group["stations"]),
                "productCount": len(group["products"]),
                "ongoingCount": int(group["ongoingCount"]),
                "totalHours": round(float(group["totalHours"]), 2),
                "expectedSalesLiters": round(float(group["expectedSalesLiters"]), 2),
            }
        )
    rows.sort(key=lambda row: (-row["totalHours"], -row["eventCount"], row["label"].casefold()))

    return {
        "source": _clean_text(payload.get("source")) or "email-fuel-outage-report",
        "sourceReceivedAt": _clean_text(payload.get("sourceReceivedAt")),
        "importedAt": _clean_text(payload.get("importedAt")),
        "reportDate": max(report_dates, default=""),
        "groupBy": group_by,
        "unmatchedStationCount": len(unmatched_stations),
        "totals": {
            "eventCount": len(items),
            "stationCount": len(total_stations),
            "productCount": len(total_products),
            "ongoingCount": sum(1 for item in items if is_outage_ongoing(item)),
            "totalHours": _sum(item.get("hours") for item in items),
            "expectedSalesLiters": _sum(item.get("expectedSalesLiters") for item in items),
        },
        "rows": rows,
    }


def replace_outage_snapshot(
    xlsx_content: bytes,
    *,
    source_message_id: str = "",
    source_received_at: str = "",
    source_email_from: str = "",
    now: datetime | None = None,
) -> dict[str, Any]:
    rows = rows_from_xlsx(xlsx_content)
    canonical_xlsx = xlsx_from_rows(rows)
    checksum = hashlib.sha256(
        json.dumps(rows, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest()
    current = _read_snapshot()
    unchanged = bool(
        current
        and current.get("checksum") == checksum
        and current.get("sourceMessageId", "") == source_message_id
    )
    if unchanged:
        return {**current, "unchanged": True}

    resolved_now = now or datetime.now(timezone.utc)
    imported_at = resolved_now.astimezone(timezone.utc).isoformat(timespec="seconds")
    station_count = len({row["ksss"] for row in rows if row.get("ksss")})
    active_count = sum(1 for row in rows if is_outage_ongoing(row))
    payload = {
        "source": "email-fuel-outage-report",
        "sourceMessageId": _clean_text(source_message_id)[:500],
        "sourceReceivedAt": _clean_text(source_received_at)[:80],
        "sourceEmailFrom": _clean_text(source_email_from).casefold()[:254],
        "importedAt": imported_at,
        "checksum": checksum,
        "rowCount": len(rows),
        "stationCount": station_count,
        "activeCount": active_count,
        "totalHours": _sum(row.get("hours") for row in rows),
        "expectedSalesLiters": _sum(row.get("expectedSalesLiters") for row in rows),
        "items": rows,
    }
    _atomic_write(outage_xlsx_path(), canonical_xlsx)
    _atomic_write(outage_snapshot_path(), json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8"))
    return {**payload, "unchanged": False}


def outage_health(now: datetime | None = None) -> dict[str, Any]:
    payload = _read_snapshot()
    if not payload:
        return {"ok": True, "rows": 0, "stations": 0, "active": 0, "importedAt": "", "stale": True}
    resolved_now = now or datetime.now(timezone.utc)
    try:
        imported_at = datetime.fromisoformat(str(payload.get("importedAt", "")).replace("Z", "+00:00"))
        stale = (resolved_now.astimezone(timezone.utc) - imported_at.astimezone(timezone.utc)).total_seconds() > stale_after_seconds()
    except (ValueError, TypeError):
        stale = True
    return {
        "ok": True,
        "rows": int(payload.get("rowCount", 0)),
        "stations": int(payload.get("stationCount", 0)),
        "active": sum(1 for item in payload.get("items") or [] if is_outage_ongoing(item)),
        "importedAt": str(payload.get("importedAt", "")),
        "sourceReceivedAt": str(payload.get("sourceReceivedAt", "")),
        "stale": stale,
    }


def get_outage_snapshot() -> dict[str, Any] | None:
    return _read_snapshot()
