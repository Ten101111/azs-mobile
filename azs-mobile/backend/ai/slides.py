"""ИИ-22, этап 1. Презентация PPTX из сохранённого ответа ИИ.

Файл собирает детерминированный построитель на python-pptx, модель в сборке не
участвует. Числа на слайдах — только из результатов: графики и таблицы строятся из
рядов и строк, которые сохранены с ответом (их подставил код из rN/pN, ИИ-17), итоги
сравнения периодов и отклонения KPI посчитал код. Текст выводов — тот, что уже
прошёл сверку чисел с результатами (ИИ-25): пункт с неподтверждённым числом снят до
сохранения ответа, а здесь текст не переписывается.

Состав (не больше 15 слайдов): титул → «Главное» (вывод и ключевые показатели) →
графики (нативные диаграммы PPTX: в PowerPoint открываются как «Изменить данные») →
таблицы (таблицы PPTX) → выводы с типом каждого пункта → рекомендации (ИИ-16) →
ограничения → источники и дата формирования. Что не поместилось в 15 слайдов,
названо на слайде ограничений.

Собрать можно только свой ответ — владение проверяет `dialogs.message()`, как и у
выгрузок ИИ-12. SQL в презентацию не попадает никогда: он есть в выгрузке Excel
для администратора.

Шаблон (решение Р-11): официальный шаблон ОНПО (.potx или .pptx), если он задан в
AI_PPTX_TEMPLATE; иначе стиль проекта (CLAUDE.md): 16:9, красная полоса 5 px сверху,
Tahoma, подвал «LUKOIL ОНПО · Инструмент АУП · Конфиденциально».

Дальше (этап 2): модель предлагает структуру слайдов (тип, заголовок, ссылки на
результаты), пользователь правит её как план ИИ-23, а собирает этот же построитель.
"""
from __future__ import annotations

import io
import math
import os
import re
import zipfile
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import (XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE,
                             XL_TICK_LABEL_POSITION)
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from .export import DEPTH_TITLES, FOOTER, GRIF, MSK, TASK_TITLES, _period, _stamp

MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
MAX_SLIDES = 15
TABLE_ROWS = 12          # строк таблицы на слайде; полностью — в выгрузке Excel (ИИ-12)
TABLE_COLS = 8
KPI_CARDS = 8
FONT = "Tahoma"

RED = RGBColor(0xC0, 0x00, 0x00)
INK = RGBColor(0x11, 0x11, 0x11)
MID = RGBColor(0x59, 0x59, 0x59)
SOFT = RGBColor(0x8C, 0x8C, 0x8C)
PALE = RGBColor(0xF2, 0xF2, 0xF2)
PINK = RGBColor(0xFB, 0xED, 0xEF)
BORDER = RGBColor(0xD9, 0xD9, 0xD9)
GRID = RGBColor(0xE6, 0xE6, 0xE6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
# Ряды графиков — как в приложении (ИИ-17): первый красный, остальные серые и различаются штрихом.
SERIES = [RED, RGBColor(0x59, 0x59, 0x59), RGBColor(0xA6, 0xA6, 0xA6), RGBColor(0x26, 0x26, 0x26),
          RGBColor(0x7F, 0x7F, 0x7F), RGBColor(0xBF, 0xBF, 0xBF)]
DASHES = [None, None, MSO_LINE_DASH_STYLE.DASH, None, MSO_LINE_DASH_STYLE.ROUND_DOT, MSO_LINE_DASH_STYLE.DASH_DOT]

TAG_W = (1.08, 1.5)      # ширина метки типа: обычная и «Рекомендация»
SECTIONS = (("happened", "Что произошло"), ("why", "Почему"), ("where", "Где именно"))
CLAIM_TITLES = {"fact": "Факт", "calc": "Расчёт", "hypothesis": "Гипотеза", "recommendation": "Рекомендация"}
DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
SPACES = re.compile(r"[^\S\u00a0\u202f]+")      # пробелы схлопываются, неразрывные (разряды чисел) — нет


# Служебные пометки ответа, которые на слайде «Ограничения» — шум: область данных уже на титуле и в
# источниках, лимит строк — в сноске таблицы, файлы и память папки — на слайде источников.
TECH_NOTES = ("Учтены файлы", "Учтена память", "Область данных:", "Подставлен фильтр области данных",
              "Ограничение вывода:", "Справочник признаков")


class NotPresentable(ValueError):
    """Из ответа нечего собрать: ошибка, план ещё не выполнен, нет ни вывода, ни данных."""


# --- содержание: какие слайды и что на них ------------------------------------------

def _clean(value) -> str:
    if isinstance(value, dict):
        value = value.get("text") or value.get("action") or ""
    return SPACES.sub(" ", str(value or "")).strip()


def _sentence(text: str) -> str:
    """Заголовок слайда с заглавной буквы: назначение шага модель пишет со строчной."""
    text = _clean(text)
    return text[:1].upper() + text[1:]


def _cut(text: str, limit: int) -> str:
    text = _clean(text)
    return text if len(text) <= limit else text[: limit - 1].rstrip(" ,;.") + "…"


def _items(values) -> list[str]:
    return [t for t in (_clean(v) for v in values or []) if t]


def chart_source(chart: dict) -> str:
    """Откуда график — теми же словами, что подпись в приложении (ИИ-17)."""
    if chart.get("auto"):
        return "тот же запрос к витрине, что и таблица ответа"
    kind, title = chart.get("sourceKind"), _clean(chart.get("sourceTitle"))
    if kind in ("file", "file_text"):
        return title or "файл пользователя"
    what = "расчёт" if kind == "python" else "запрос к витрине"
    source = chart.get("source") or ""
    return f"{what} «{title}» ({source})" if title else f"{what} {source}".strip()


def _caption(chart: dict, answer: dict) -> str:
    plan = answer.get("plan") or {}
    parts = [f"Источник: {chart_source(chart)}"]
    period = chart.get("period") or plan.get("period")
    if period:
        parts.append(f"Период: {_clean(period)}")
    filters = [_clean(f) for f in plan.get("filters") or [] if _clean(f)]
    if filters:
        parts.append("Фильтры: " + "; ".join(filters))
    if answer.get("scopeLabel"):
        parts.append(f"Область: {_clean(answer['scopeLabel'])}")
    return " · ".join(parts)


def _step(answer: dict, source: str) -> dict | None:
    return next((s for s in answer.get("steps") or []
                 if source and str(s.get("resultId") or "").lower() == str(source).lower()), None)


def _points(analysis: dict) -> list[dict]:
    """Пункты выводов с типом (ИИ-25): факт, расчёт, гипотеза; формула или что проверить."""
    marks = analysis.get("claims") or {}
    out = []
    for key, title in SECTIONS:
        for index, text in enumerate(_items(analysis.get(key))):
            claim = (marks.get(key) or [])[index] if index < len(marks.get(key) or []) else {}
            kind = (claim or {}).get("type") or ""
            detail = ""
            if kind == "calc" and claim.get("formula"):
                detail = "Расчёт: " + "; ".join(_clean(f) for f in claim["formula"][:2])
            elif kind == "hypothesis" and claim.get("check"):
                detail = "Проверить: " + _clean(claim["check"])
            out.append({"section": title, "type": kind, "text": _cut(text, 420), "detail": _cut(detail, 260)})
    return out


def _recommendations(analysis: dict) -> list[dict]:
    recs = analysis.get("recommendations")
    if isinstance(recs, list):
        return [{"action": _cut(r.get("action"), 300), "basis": _cut(r.get("basis"), 240),
                 "effect": _cut(r.get("effect"), 200), "limits": _cut(r.get("limits"), 200),
                 "confidence": _clean(r.get("confidence")), "source": _cut(r.get("source"), 160)}
                for r in recs if _clean(r.get("action"))]
    # Ответы до ИИ-16: только текст действия.
    return [{"action": _cut(a, 300)} for a in _items(analysis.get("actions"))]


def _table(title: str, columns: list, rows: list, *, source: str, truncated: bool, note: str = "") -> dict:
    return {"kind": "table", "title": _sentence(_cut(title, 120)) or "Таблица", "columns": [str(c) for c in columns],
            "rows": [list(r) for r in rows], "source": source, "truncated": truncated, "note": note}


def _headline_size(text: str, width_in: float = 12.13 - 0.35) -> tuple[float, float]:
    """Кегль и высота главного вывода: длинный вывод — мельче, но без переполнения."""
    size = 30 if len(text) <= 90 else (25 if len(text) <= 180 else 20)
    return size, _lines(text, width_in, size, bold=True) * size * 1.15 / 72


def _kpi_block(chart: dict) -> float:
    """Высота блока показателей на «Главном»: подзаголовок, одна или две строки карточек, подпись источника."""
    count = len((chart.get("cards") or [])[:KPI_CARDS])
    return 0.42 + (1.75 if count <= 4 else 3.45) + 0.35


def _fit_points(points: list[dict], room: float, width_in: float = 12.13) -> int:
    """Сколько пунктов «Что произошло» встаёт на слайд «Главное» под выводом (с подзаголовком)."""
    used, count = 0.42, 0
    for point in points[:5]:
        used += _point_height(point, width_in)
        if used > room:
            break
        count += 1
    return count


def outline(message: dict, meta: dict) -> list[dict]:
    """Слайды по порядку. meta: role, scope, source, tables, catalog, now, email — см. api."""
    answer = message.get("answer") or {}
    card = answer.get("planCard") or {}
    if not answer.get("ok") or (card and card.get("status") not in (None, "done")):
        raise NotPresentable("Из этого ответа презентацию не собрать: в нём нет готового результата")
    analysis = answer.get("analysis") or {}
    headline = _clean(analysis.get("headline")) or _clean(answer.get("summary"))
    charts = [c for c in answer.get("charts") or [] if c.get("type")]
    rows = answer.get("rows") or []
    if not (headline or charts or rows):
        raise NotPresentable("В ответе нет ни вывода, ни данных для слайдов")

    question = _clean(message.get("question") or answer.get("question"))
    depth = answer.get("depth") or "fast"
    task = answer.get("taskType") or "lookup"
    plan = answer.get("plan") or {}
    frame = answer.get("frame") or {}
    period = _clean(plan.get("period") or frame.get("period")) or _period(frame, answer.get("sql") or "")
    level = f"Уровень «{DEPTH_TITLES.get(depth, depth)}» · задача — {TASK_TITLES.get(task, task)}"
    scope = _clean(answer.get("scopeLabel") or meta.get("scope")) or "—"

    kpi = next((c for c in charts if c["type"] == "kpi" and c.get("cards")), None)
    _size, head_h = _headline_size(_cut(headline, 420))
    if kpi and 0.68 + head_h + 0.45 + _kpi_block(kpi) > BODY_BOTTOM:
        kpi = None                       # длинный вывод и много карточек: показатели — отдельным слайдом
    cards = (kpi or {}).get("cards", [])[:KPI_CARDS]
    head = {"kind": "headline", "text": _cut(headline, 420) or "Результат запроса — на следующих слайдах",
            "cards": cards, "kpi": kpi, "caption": _caption(kpi, answer) if kpi else "", "points": []}
    title = {"kind": "title", "question": question, "level": level, "scope": scope, "period": period,
             "stamp": _stamp(message.get("createdAt"))}

    visuals: list[dict] = []
    for chart in charts:
        if chart is kpi:
            continue
        if chart["type"] == "kpi":
            if chart.get("cards"):
                visuals.append({"kind": "kpi", "chart": chart, "caption": _caption(chart, answer),
                                "note": _clean(chart.get("note")), "headline": ""})
        elif chart["type"] == "table":
            visuals.append(_table(chart.get("title") or "Таблица", chart.get("columns") or [], chart.get("rows") or [],
                                  source=_caption(chart, answer), truncated=False, note=_clean(chart.get("note"))))
        else:
            visuals.append({"kind": "chart", "chart": chart, "caption": _caption(chart, answer),
                            "note": _clean(chart.get("note")), "headline": ""})
    tables: list[dict] = []
    auto = next((c for c in charts if c.get("auto")), None)
    # Простой график «Лёгкого» уже подписан значениями всех строк: та же таблица из двух столбцов его повторяла бы.
    duplicate = auto is not None and len(answer.get("columns") or []) <= 2 and len(rows) == len(auto.get("x") or [])
    if rows and not duplicate:
        main = str(frame.get("mainResult") or ("r1" if depth == "fast" else ""))
        step = _step(answer, main) or {}
        tables.append(_table(step.get("purpose") or ("Результат запроса" if depth == "fast" else "Данные расчёта"),
                             answer.get("columns") or [], rows, truncated=bool(answer.get("truncated")),
                             source=f"Источник: запрос к витрине{(' ' + main) if main else ''} · Область: {scope}"))
    for extra in answer.get("tables") or []:
        if extra.get("rows"):
            what = "расчёт" if str(extra.get("id") or "").lower().startswith("p") else "запрос к витрине"
            tables.append(_table(extra.get("title") or extra.get("id") or "Таблица", extra.get("columns") or [],
                                 extra["rows"], truncated=bool(extra.get("truncated")),
                                 source=" ".join(f"Источник: {what} {extra.get('id') or ''}".split()) + f" · Область: {scope}"))

    points = _points(analysis)
    # «Главное»: вывод, ключевые показатели и первые пункты «Что произошло» — сколько поместится.
    _size, head_h = _headline_size(head["text"])
    room = BODY_BOTTOM - (0.68 + head_h + 0.45)
    if cards:
        room -= _kpi_block(kpi) + 0.3
    happened = [p for p in points if p["section"] == SECTIONS[0][1]]
    fitted = _fit_points(happened, room) if room > 0.9 else 0
    head["points"] = happened[:fitted]
    points = [p for p in points if p not in head["points"]]

    recs = _recommendations(analysis)
    limits = _items(analysis.get("limitations"))
    for note in _items(answer.get("notes")):
        if note not in limits and not note.startswith(TECH_NOTES):
            limits.append(note)
    if analysis.get("recommendationsWithheld") and not recs:
        limits.append("Рекомендаций нет: для них недостаточно данных (правила ИИ-16).")

    text_slides = _paginate_points(points, pages=2) + _paginate_recs(recs, analysis, pages=2)
    # Одна строка вывода на пустом слайде — пустой слайд: вывод становится заголовком первого графика или таблицы.
    merge = not cards and not head["points"] and bool(visuals or tables)
    fixed = 1 + (0 if merge else 1) + len(text_slides) + 1 + 1   # титул, главное, …, ограничения, источники
    room_slides = max(1, MAX_SLIDES - fixed)
    shown_visuals = visuals[:room_slides]
    shown_tables = tables[: max(0, room_slides - len(shown_visuals))]
    if len(shown_visuals) < len(visuals) or len(shown_tables) < len(tables):
        limits.append(f"В презентацию вошли графиков: {sum(v['kind'] == 'chart' for v in shown_visuals)} из "
                      f"{sum(v['kind'] == 'chart' for v in visuals)}, таблиц: "
                      f"{len(shown_tables) + sum(v['kind'] == 'table' for v in shown_visuals)} из "
                      f"{len(tables) + sum(v['kind'] == 'table' for v in visuals)} — не больше {MAX_SLIDES} слайдов; "
                      "остальное — в ответе и в выгрузке Excel.")
    for table in shown_tables + [v for v in shown_visuals if v["kind"] == "table"]:
        if table["truncated"]:
            limits.append(f"Таблица «{table['title']}»: результат запроса обрезан лимитом строк — в витрине строк больше.")
    for visual in shown_visuals:
        floor = _waterfall_floor(visual["chart"]) if visual["kind"] == "chart" else None
        if floor is not None:
            visual["note"] = " ".join(filter(None, [visual["note"], f"Ось значений начинается с "
                                                    f"{fmt_number(floor)}, а не с нуля: так видны изменения."]))

    body = shown_visuals + shown_tables
    if merge:
        body[0] = dict(body[0], headline=head["text"])
        slides = [title] + body + text_slides
    else:
        slides = [title, head] + body + text_slides
    if limits:
        slides.append({"kind": "limits", "items": [_cut(t, 300) for t in limits[:10]]})
    slides.append({"kind": "sources", "rows": _sources(message, meta, period, level)})
    return slides


def _sources(message: dict, meta: dict, period: str, level: str) -> list[tuple[str, str]]:
    """Последний слайд: откуда цифры (как паспорт выгрузки ИИ-12, без SQL)."""
    answer = message.get("answer") or {}
    plan, frame = answer.get("plan") or {}, answer.get("frame") or {}
    results = []
    for step in answer.get("steps") or []:
        rid = step.get("resultId")
        if not rid or step.get("ok") is False:
            continue
        what = {"python": "расчёт в песочнице", "file": "файл", "file_text": "файл"}.get(step.get("kind"), "запрос к витрине")
        rows = step.get("rows")
        results.append(f"{rid} — {_clean(step.get('purpose') or step.get('label')) or what} ({what}"
                       + (f", строк: {rows}" if isinstance(rows, int) else "") + ")")
    if not results and answer.get("rows"):
        results.append(f"r1 — основной запрос ответа (запрос к витрине, строк: {len(answer['rows'])})")
    files = [_clean(f.get("name")) for f in (answer.get("context") or {}).get("files") or [] if f.get("name")]
    filters = [_clean(f) for f in (plan.get("filters") or frame.get("filters") or []) if _clean(f)]
    tables = meta.get("tables") or []
    now = meta.get("now") or datetime.now(MSK)
    return [
        ("Вопрос", _cut(message.get("question") or answer.get("question"), 300)),
        ("Источник данных", meta.get("source", "—") + (f": {', '.join(tables)}" if tables else "")),
        ("Результаты расчёта", _cut("; ".join(results), 420) or "—"),
        ("Период", period or "—"),
        ("Фильтры и параметры", "; ".join([level] + filters)),
        ("Файлы-источники", _cut(", ".join(files), 300) if files else "не использовались"),
        ("Роль и область данных", f"{meta.get('role') or '—'}; область: {_clean(answer.get('scopeLabel')) or '—'} "
                                  "— подставлена системой, а не выбрана моделью"),
        ("Ответ сформирован", _stamp(message.get("createdAt"))),
        ("Презентация сформирована", now.astimezone(MSK).strftime("%d.%m.%Y %H:%M МСК")
         + (f", {meta['email']}" if meta.get("email") else "")),
        ("Номер ответа", f"сообщение {message.get('id')}" + (f", запись журнала {message.get('journalId')}"
                                                            if message.get("journalId") else "")),
        ("Версия каталога", meta.get("catalog") or "—"),
        ("Гриф", f"{GRIF}. Не пересылайте за пределы компании."),
    ]


# --- разбивка текста по слайдам -------------------------------------------------------
# Высоты оцениваются по числу знаков: python-pptx не измеряет текст, а переполнение
# слайда хуже лишнего слайда.

BODY_TOP, BODY_BOTTOM = 1.55, 6.72


def _lines(text: str, width_in: float, size_pt: float, bold: bool = False) -> int:
    """Строк текста по числу знаков — с запасом: у Tahoma средний знак ≈ 0,55 кегля, у жирного шире."""
    per_line = max(10, int(width_in * 72 / (size_pt * (0.68 if bold else 0.62))))
    return max(1, sum(math.ceil(max(1, len(part)) / per_line) for part in str(text).split("\n")))


def _point_height(point: dict, width_in: float) -> float:
    text_h = _lines(point["text"], width_in - 1.3, 14) * 14 * 1.25 / 72
    detail_h = _lines(point["detail"], width_in - 1.3, 11) * 11 * 1.25 / 72 + 0.04 if point["detail"] else 0
    return max(0.3, text_h) + detail_h + 0.16


def _paginate_points(points: list[dict], pages: int, width_in: float = 12.13) -> list[dict]:
    if not points:
        return []
    slides, current, used, section = [], [], 0.0, None
    room = BODY_BOTTOM - BODY_TOP - 0.3
    for point in points:
        extra = 0.42 if point["section"] != section else 0
        height = extra + _point_height(point, width_in)
        if current and used + height > room:
            slides.append(current)
            current, used, section = [], 0.0, None
            extra = 0.42
            height = extra + _point_height(point, width_in)
        current.append(point)
        used += height
        section = point["section"]
    slides.append(current)
    kept = slides[:pages]
    dropped = sum(len(s) for s in slides[pages:])
    out = [{"kind": "points", "points": page, "part": i + 1, "parts": len(kept)} for i, page in enumerate(kept)]
    if dropped:
        out[-1]["more"] = f"Ещё пунктов: {dropped} — в ответе ИИ."
    return out


REC_INDENT = 1.95         # отступ текста рекомендации: полоса и метка «Рекомендация» слева


def _rec_height(rec: dict, width_in: float) -> float:
    h = _lines(rec["action"], width_in - REC_INDENT, 15, bold=True) * 15 * 1.25 / 72
    for key in ("basis", "effect", "limits"):
        if rec.get(key):
            h += _lines(rec[key] + " Ожидаемый эффект: ", width_in - REC_INDENT, 11) * 11 * 1.3 / 72
    if rec.get("confidence") or rec.get("source"):
        h += 11 * 1.3 / 72
    return h + 0.42


def _paginate_recs(recs: list[dict], analysis: dict, pages: int, width_in: float = 12.13) -> list[dict]:
    if not recs:
        return []
    note = _clean(analysis.get("recommendationNote"))
    room = BODY_BOTTOM - BODY_TOP - (0.45 if note else 0)
    slides, current, used = [], [], 0.0
    for rec in recs:
        height = _rec_height(rec, width_in)
        if current and used + height > room:
            slides.append(current)
            current, used = [], 0.0
        current.append(rec)
        used += height
    slides.append(current)
    kept = slides[:pages]
    out = [{"kind": "recs", "recs": page, "note": note, "part": i + 1, "parts": len(kept)} for i, page in enumerate(kept)]
    dropped = sum(len(s) for s in slides[pages:])
    if dropped:
        out[-1]["more"] = f"Ещё рекомендаций: {dropped} — в ответе ИИ."
    return out


# --- числа --------------------------------------------------------------------------

def _is_number(value) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool) and not (
        isinstance(value, float) and (math.isnan(value) or math.isinf(value)))


def _places(values) -> int:
    places = 0
    for v in values:
        if isinstance(v, float) and _is_number(v) and not v.is_integer():
            text = f"{v:.6f}".rstrip("0")
            places = max(places, len(text.split(".")[1]) if "." in text else 0)
    return min(places, 3)


def fmt_number(value, places: int | None = None, signed: bool = False) -> str:
    """Русская запись числа: пробел между разрядами, запятая, длинный минус."""
    if not _is_number(value):
        return "—" if value is None else str(value)
    if places is None:
        places = _places([value])
    text = f"{abs(value):,.{places}f}".replace(",", " ").replace(".", ",")
    if value < 0 and text.strip("0, "):
        return "−" + text
    return ("+" + text) if signed and value > 0 else text


def _cell(value, places: int | None) -> str:
    if isinstance(value, str) and DATE_RE.match(value):
        return f"{value[8:10]}.{value[5:7]}.{value[:4]}"
    if isinstance(value, bool):
        return "да" if value else "нет"
    if _is_number(value):
        return fmt_number(value, places)
    return _cut(value, 60) if value not in (None, "") else "—"


def _excel_format(places: int) -> str:
    return "#,##0" if places == 0 else "#,##0." + "0" * places


def _with_unit(text: str, unit: str) -> str:
    return f"{text} {unit}" if unit and unit not in text else text


# --- оформление ---------------------------------------------------------------------

class Deck:
    """Геометрия и примитивы: всё в дюймах от размеров слайда (шаблон может быть не 16:9)."""

    def __init__(self, template: str | None = None):
        self.prs, self.templated = _base(template)
        self.W = self.prs.slide_width / 914400
        self.H = self.prs.slide_height / 914400
        self.L = 0.6
        self.CW = self.W - 2 * self.L
        self.layout = _blank_layout(self.prs)
        self.total = 0

    # примитивы
    def slide(self):
        slide = self.prs.slides.add_slide(self.layout)
        for shape in list(slide.placeholders):
            shape._element.getparent().remove(shape._element)
        return slide

    def text(self, slide, left, top, width, height, runs, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, spacing=1.1, kicker=False):
        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        frame = box.text_frame
        frame.word_wrap = True
        frame.auto_size = None
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = anchor
        paragraphs = runs if isinstance(runs, list) and runs and isinstance(runs[0], list) else [runs]
        for index, para_runs in enumerate(paragraphs):
            para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            para.alignment = align
            para.line_spacing = spacing
            if isinstance(para_runs, str):
                para_runs = [(para_runs, {})]
            for chunk, style in para_runs:
                run = para.add_run()
                run.text = chunk
                font = run.font
                font.name = FONT
                font.size = Pt(style.get("size", size))
                font.bold = style.get("bold", bold)
                font.italic = style.get("italic", False)
                font.color.rgb = style.get("color", color)
                if kicker:
                    run._r.get_or_add_rPr().set("spc", "80")
        return box

    def rect(self, slide, left, top, width, height, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, weight=0.75):
        box = slide.shapes.add_shape(shape, Inches(left), Inches(top), Inches(width), Inches(height))
        if fill is None:
            box.fill.background()
        else:
            box.fill.solid()
            box.fill.fore_color.rgb = fill
        if line is None:
            box.line.fill.background()
        else:
            box.line.color.rgb = line
            box.line.width = Pt(weight)
        _no_style(box)
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            box.adjustments[0] = 0.18
        return box

    def line(self, slide, x1, y1, x2, y2, color=BORDER, weight=0.75):
        conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        conn.line.color.rgb = color
        conn.line.width = Pt(weight)
        _no_style(conn)
        return conn

    def frame(self, slide, number: int, kicker: str = "", title: str = "") -> float:
        """Постоянные элементы: полоса, надзаголовок, заголовок, подвал с грифом и номером. Возвращает верх тела."""
        if not self.templated:
            self.rect(slide, 0, 0, self.W, 5 / 96, fill=RED)
        if kicker:
            self.text(slide, self.L, 0.34, self.CW - 2.4, 0.26, kicker.upper(), size=10.5, color=RED, bold=True,
                      kicker=True)
        top = BODY_TOP
        if title:
            size = 24 if len(title) <= 70 else (21 if len(title) <= 110 else 18)
            height = _lines(title, self.CW, size, bold=True) * size * 1.12 / 72
            self.text(slide, self.L, 0.64, self.CW, height + 0.1, title, size=size, bold=True, anchor=MSO_ANCHOR.TOP,
                      spacing=1.0)
            top = max(BODY_TOP, 0.64 + height + 0.3)
        y = self.H - 0.5
        self.line(slide, self.L, y, self.W - self.L, y)
        self.text(slide, self.L, y + 0.1, self.CW - 1.2, 0.24, FOOTER, size=9, color=SOFT)
        self.text(slide, self.W - self.L - 1.2, y + 0.1, 1.2, 0.24, f"{number} / {self.total}", size=9, color=SOFT,
                  align=PP_ALIGN.RIGHT)
        return top

    def tag(self, slide, left, top, kind: str):
        title = CLAIM_TITLES.get(kind)
        if not title:
            return
        fill, color = {"fact": (INK, WHITE), "calc": (MID, WHITE), "hypothesis": (PINK, RED),
                       "recommendation": (RED, WHITE)}.get(kind, (PALE, INK))
        box = self.rect(slide, left, top, TAG_W[kind == "recommendation"], 0.27, fill=fill,
                        shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        frame = box.text_frame
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        para = frame.paragraphs[0]
        para.alignment = PP_ALIGN.CENTER
        run = para.add_run()
        run.text = title.upper()
        run.font.name, run.font.size, run.font.bold, run.font.color.rgb = FONT, Pt(8.5), True, color
        run._r.get_or_add_rPr().set("spc", "40")

    def save(self) -> bytes:
        buffer = io.BytesIO()
        self.prs.save(buffer)
        return buffer.getvalue()


def _no_style(shape) -> None:
    """Убрать ссылку на стиль темы: иначе фигура наследует тень и обводку темы."""
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


def _base(template: str | None):
    path = Path(template) if template else None
    if path and path.is_file():
        data = path.read_bytes()
        if path.suffix.lower() == ".potx":
            data = _potx_as_pptx(data)
        prs = Presentation(io.BytesIO(data))
        _drop_slides(prs)
        return prs, True
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    return prs, False


def _potx_as_pptx(data: bytes) -> bytes:
    """Шаблон .potx открывается как презентация: меняется только тип главной части."""
    source, out = zipfile.ZipFile(io.BytesIO(data)), io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            blob = source.read(item.filename)
            if item.filename == "[Content_Types].xml":
                blob = blob.replace(b"presentationml.template.main+xml", b"presentationml.presentation.main+xml")
            target.writestr(item, blob)
    return out.getvalue()


def _drop_slides(prs) -> None:
    ids = prs.slides._sldIdLst
    for sld in list(ids):
        prs.part.drop_rel(sld.get(qn("r:id")))
        ids.remove(sld)


def _blank_layout(prs):
    names = ("blank", "пустой")
    for layout in prs.slide_layouts:
        if any(n in (layout.name or "").lower() for n in names):
            return layout
    return min(prs.slide_layouts, key=lambda layout: len(layout.placeholders))


# --- слайды -------------------------------------------------------------------------

def _title_slide(deck: Deck, spec: dict, number: int):
    s = deck.slide()
    deck.frame(s, number)
    L, W = deck.L, deck.W
    badge = deck.rect(s, W - L - 2.6, 0.34, 2.6, 0.34, line=RED, shape=MSO_SHAPE.ROUNDED_RECTANGLE, weight=1)
    frame = badge.text_frame
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = frame.paragraphs[0]
    para.alignment = PP_ALIGN.CENTER
    run = para.add_run()
    run.text = GRIF.upper()
    run.font.name, run.font.size, run.font.bold, run.font.color.rgb = FONT, Pt(10), True, RED
    run._r.get_or_add_rPr().set("spc", "60")

    question = _cut(spec["question"], 220) or "Ответ ИИ-аналитика"
    size = 32 if len(question) <= 80 else (26 if len(question) <= 150 else 22)
    height = _lines(question, deck.CW - 1.0, size, bold=True) * size * 1.12 / 72
    top = 1.75
    deck.text(s, L + 0.3, top - 0.46, deck.CW - 0.3, 0.3, "ИИ-АНАЛИТИК · ОТВЕТ НА ВОПРОС", size=11, color=RED,
              bold=True, kicker=True)
    deck.rect(s, L, top + 0.06, 0.09, max(0.6, height - 0.04), fill=RED)
    deck.text(s, L + 0.3, top, deck.CW - 1.0, height + 0.2, question, size=size, bold=True, spacing=1.0)
    y = top + height + 0.55
    lines = [[("Область данных: ", {"bold": True, "color": INK}), (spec["scope"], {})],
             [(spec["level"], {})],
             [("Период: ", {"bold": True, "color": INK}), (spec["period"] or "—", {})],
             [("Ответ сформирован: ", {"bold": True, "color": INK}), (spec["stamp"], {})]]
    deck.text(s, L + 0.3, y, deck.CW - 1.0, 1.6, lines, size=14, color=MID, spacing=1.35)


def _kpi_cards(deck: Deck, slide, cards: list[dict], unit: str, top: float, height: float = 1.55):
    per_row = min(4, len(cards))
    gap = 0.25
    width = (deck.CW - gap * (per_row - 1)) / per_row
    card_h = height
    for index, card in enumerate(cards):
        row, col = divmod(index, per_row)
        left, y = deck.L + col * (width + gap), top + row * (card_h + gap)
        deck.rect(slide, left, y, width, card_h, fill=PALE)
        deck.rect(slide, left, y, 0.07, card_h, fill=RED if (card.get("delta") or 0) < 0 else MID)
        label = _cut(card.get("label"), 70)
        deck.text(slide, left + 0.25, y + 0.14, width - 0.4, 0.36, label, size=11.5, color=MID)
        value = card.get("value")
        value_text = _with_unit(fmt_number(value), unit) if _is_number(value) else _cut(value, 30)
        size = 28 if len(value_text) <= 12 else (22 if len(value_text) <= 18 else 17)
        deck.text(slide, left + 0.25, y + 0.46, width - 0.4, 0.5, value_text, size=size, bold=True)
        lines = []
        if card.get("base") is not None:
            lines.append([(f"{_clean(card.get('baseLabel')) or 'База'}: ", {}),
                          (fmt_number(card["base"]) if _is_number(card["base"]) else _clean(card["base"]), {})])
        delta, pct = card.get("delta"), card.get("deltaPct")
        if _is_number(delta) or _is_number(pct):
            down = (delta if _is_number(delta) else pct) < 0
            parts = []
            if _is_number(delta):
                parts.append(fmt_number(delta, signed=True))
            if _is_number(pct):
                parts.append(f"{fmt_number(pct, signed=True)} %")
            text = parts[0] + (f" ({parts[1]})" if len(parts) > 1 else "")
            lines.append([("Изменение: ", {}), (text, {"bold": True, "color": RED if down else INK})])
        if lines:
            lines_top = y + max(1.0, card_h - 0.16 - 0.23 * len(lines))
            deck.text(slide, left + 0.25, lines_top, width - 0.4, 0.23 * len(lines) + 0.05,
                      lines, size=10.5, color=MID, spacing=1.15)


def _kpi_slide(deck: Deck, spec: dict, number: int):
    chart = spec["chart"]
    s = deck.slide()
    name = _sentence(_cut(chart.get("title") or "Ключевые показатели", 150))
    if spec.get("headline"):
        top = deck.frame(s, number, kicker="Главное", title=spec["headline"])
        deck.text(s, deck.L, top - 0.05, deck.CW, 0.3, name, size=13, bold=True, color=MID)
        top += 0.4
    else:
        top = deck.frame(s, number, kicker="Показатели", title=name)
    cards = chart["cards"][:KPI_CARDS]
    _kpi_cards(deck, s, cards, chart.get("unit") or "", top + 0.1, height=1.75 if len(cards) <= 4 else 1.6)
    lines = [spec["caption"]] + ([spec["note"]] if spec.get("note") else [])
    deck.text(s, deck.L, deck.H - 0.62 - 0.2 * len(lines), deck.CW, 0.2 * len(lines), [[(t, {})] for t in lines],
              size=9.5, color=SOFT, spacing=1.05)


def _headline_slide(deck: Deck, spec: dict, number: int):
    s = deck.slide()
    deck.frame(s, number, kicker="Главное")
    text = spec["text"]
    size, height = _headline_size(text, deck.CW - 0.35)
    deck.rect(s, deck.L, 0.72, 0.09, height, fill=RED)
    deck.text(s, deck.L + 0.35, 0.68, deck.CW - 0.35, height + 0.2, text, size=size, bold=True, spacing=1.05)
    y = 0.68 + height + 0.45
    if spec["cards"]:
        kpi = spec["kpi"] or {}
        y += 0.42
        title = _clean(kpi.get("title"))
        if title:
            deck.text(s, deck.L, y - 0.4, deck.CW, 0.3, title, size=13, bold=True, color=MID)
        cards_h = 1.75 if len(spec["cards"]) <= 4 else 3.45
        _kpi_cards(deck, s, spec["cards"], kpi.get("unit") or "", y, height=1.75 if len(spec["cards"]) <= 4 else 1.6)
        # две строки карточек: 2 × 1,6 + зазор — как в оценке места в outline()
        y += cards_h + 0.3
        deck.text(s, deck.L, deck.H - 0.86, deck.CW, 0.26, spec["caption"], size=9.5, color=SOFT)
    if spec["points"]:
        _draw_points(deck, s, spec["points"], y)


def _draw_points(deck: Deck, s, points: list[dict], y: float) -> float:
    """Пункты с меткой типа (ИИ-25): подзаголовок раздела, метка, текст, формула или что проверить."""
    section = None
    text_w = deck.CW - 1.3
    for point in points:
        if point["section"] != section:
            section = point["section"]
            deck.text(s, deck.L, y + 0.04, deck.CW, 0.28, section.upper(), size=10.5, color=MID, bold=True, kicker=True)
            y += 0.42
        deck.tag(s, deck.L, y + 0.01, point["type"])
        text_h = _lines(point["text"], text_w, 14) * 14 * 1.25 / 72
        deck.text(s, deck.L + 1.3, y, text_w, max(0.3, text_h), point["text"], size=14, spacing=1.08)
        y += max(0.3, text_h)
        if point["detail"]:
            detail_h = _lines(point["detail"], text_w, 11) * 11 * 1.25 / 72
            deck.text(s, deck.L + 1.3, y + 0.04, text_w, detail_h, point["detail"], size=11, color=MID)
            y += detail_h + 0.04
        y += 0.16
    return y


def _chart_slide(deck: Deck, spec: dict, number: int, index: int, count: int):
    chart = spec["chart"]
    s = deck.slide()
    name = _sentence(_cut(chart.get("title") or "График", 150))
    if spec.get("headline"):
        top = deck.frame(s, number, kicker="Главное", title=spec["headline"])
        deck.text(s, deck.L, top - 0.05, deck.CW, 0.3, name, size=13, bold=True, color=MID)
        top += 0.35
    else:
        top = deck.frame(s, number, kicker=f"График {index} из {count}" if count > 1 else "График", title=name)
    footer_lines = [spec["caption"]] + ([spec["note"]] if spec["note"] else [])
    totals = _compare_totals(chart)
    bottom = deck.H - 0.62 - 0.2 * sum(_lines(t, deck.CW, 9.5) for t in footer_lines)
    if totals:
        deck.text(s, deck.L, top - 0.02, deck.CW, 0.32, totals, size=12.5, color=MID)
        top += 0.38
    _draw_chart(deck, s, chart, deck.L, top, deck.CW, bottom - top - 0.12)
    deck.text(s, deck.L, bottom, deck.CW, deck.H - 0.6 - bottom, [[(t, {})] for t in footer_lines], size=9.5,
              color=SOFT, spacing=1.05)


def _compare_totals(chart: dict):
    periods = chart.get("periods") or []
    if chart.get("type") != "compare" or not periods or any(not _is_number(p.get("total")) for p in periods):
        return None
    unit = chart.get("unit") or ""
    word = "В среднем" if chart.get("aggregate") == "avg" else "Итого"
    runs = [(f"{word}: ", {"bold": True})]
    for i, p in enumerate(periods):
        runs.append((f"{p.get('label')} — {_with_unit(fmt_number(p['total'], 1 if chart.get('aggregate') == 'avg' else None), unit)}"
                     + ("; " if i < len(periods) - 1 else ""), {"color": INK if i == 0 else MID}))
    pct = chart.get("changePct")
    if _is_number(pct):
        runs.append((f"   {fmt_number(pct, 1, signed=True)} %", {"bold": True, "color": RED if pct < 0 else INK}))
    return runs


def _draw_chart(deck: Deck, slide, chart: dict, left, top, width, height):
    kind = chart.get("type")
    if kind == "scatter":
        return _scatter(slide, chart, left, top, width, height)
    if kind == "waterfall":
        drawn = _waterfall(slide, chart, left, top, width, height)
        if drawn is not None:
            return drawn
    xs = [str(x) for x in chart.get("x") or []]
    series = [s for s in chart.get("series") or [] if s.get("values")]
    if not xs or not series:
        return None
    values = [v for s in series for v in s["values"] if _is_number(v)]
    places = min(_places(values), 2)
    unit = chart.get("unit") or ""
    data = CategoryChartData(number_format=_excel_format(places))
    horizontal = kind == "bar" and (max(len(x) for x in xs) > 12 or len(xs) > 14) and len(xs) <= 30
    order = list(reversed(range(len(xs)))) if horizontal else list(range(len(xs)))
    data.categories = [xs[i] for i in order]
    for n, s in enumerate(series, start=1):
        vals = list(s["values"]) + [None] * (len(xs) - len(s["values"]))
        data.add_series(_clean(s.get("name")) or f"Ряд {n}",
                        [vals[i] if _is_number(vals[i]) else None for i in order])
    chart_type = {"line": XL_CHART_TYPE.LINE, "compare": XL_CHART_TYPE.LINE, "stacked_bar": XL_CHART_TYPE.COLUMN_STACKED,
                  "waterfall": XL_CHART_TYPE.COLUMN_CLUSTERED}.get(kind, XL_CHART_TYPE.COLUMN_CLUSTERED)
    if horizontal:
        chart_type = XL_CHART_TYPE.BAR_CLUSTERED
    frame = slide.shapes.add_chart(chart_type, Inches(left), Inches(top), Inches(width), Inches(height), data)
    ch = frame.chart
    _style_axes(ch, places, unit, chart)
    if any(v < 0 for v in values):
        ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    plot = ch.plots[0]
    lines = chart_type == XL_CHART_TYPE.LINE
    many = len(xs) > 14
    for i, ser in enumerate(plot.series):
        color = SERIES[i % len(SERIES)]
        dash = DASHES[i % len(DASHES)]
        if kind == "compare":
            color, dash = (RED, None) if i == 0 else (RGBColor(0x7F, 0x7F, 0x7F), MSO_LINE_DASH_STYLE.DASH)
        if lines:
            ser.smooth = False
            ser.format.line.color.rgb = color
            ser.format.line.width = Pt(2.5 if i == 0 else 1.75)
            if dash is not None:
                ser.format.line.dash_style = dash
            ser.marker.style = XL_MARKER_STYLE.NONE if many else XL_MARKER_STYLE.CIRCLE
            if not many:
                ser.marker.size = 6
                ser.marker.format.fill.solid()
                ser.marker.format.fill.fore_color.rgb = color
                ser.marker.format.line.color.rgb = color
        else:
            ser.invert_if_negative = False
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = color
    if not lines:
        plot.gap_width = 70 if len(xs) <= 12 else 40
        if chart_type == XL_CHART_TYPE.COLUMN_STACKED:
            plot.overlap = 100
        elif len(series) > 1:
            plot.overlap = -8
    if len(series) == 1 and not lines:
        # Один ряд: при отрицательных значениях снижение красным, рост серым — как в приложении.
        raw = [series[0]["values"][i] if i < len(series[0]["values"]) else None for i in order]
        if any(_is_number(v) and v < 0 for v in raw):
            for point_index, v in enumerate(raw):
                point = plot.series[0].points[point_index]
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = RED if _is_number(v) and v < 0 else MID
        if len(xs) <= 16:
            plot.has_data_labels = True
            labels = plot.data_labels
            labels.number_format = _excel_format(places)
            labels.number_format_is_linked = False
            labels.font.size = Pt(9.5)
            labels.font.name = FONT
            labels.font.color.rgb = INK
            labels.position = XL_LABEL_POSITION.OUTSIDE_END
    ch.has_legend = len(series) > 1
    if ch.has_legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(10)
        ch.legend.font.name = FONT
    return frame


def _style_axes(ch, places: int, unit: str, chart: dict, xy: bool = False):
    ch.has_title = False                 # название — заголовком слайда; PowerPoint иначе ставит имя ряда
    ch.font.name = FONT
    ch.font.size = Pt(10)
    ch.font.color.rgb = MID
    value_axis = ch.value_axis
    value_axis.has_major_gridlines = True
    value_axis.major_gridlines.format.line.color.rgb = GRID
    value_axis.major_gridlines.format.line.width = Pt(0.75)
    value_axis.format.line.fill.background()
    value_axis.tick_labels.number_format = _excel_format(places)
    value_axis.tick_labels.number_format_is_linked = False
    value_axis.tick_labels.font.size = Pt(9.5)
    title = _clean(chart.get("yTitle")) if xy else unit
    if title:
        value_axis.has_title = True
        value_axis.axis_title.text_frame.text = title
        _axis_title_font(value_axis)
    category_axis = ch.category_axis
    category_axis.format.line.color.rgb = BORDER
    category_axis.tick_labels.font.size = Pt(9.5)
    category_axis.has_major_gridlines = False
    x_title = _clean(chart.get("xTitle"))
    if xy and x_title:
        category_axis.has_title = True
        category_axis.axis_title.text_frame.text = x_title
        _axis_title_font(category_axis)


def _axis_title_font(axis):
    for para in axis.axis_title.text_frame.paragraphs:
        for run in para.runs:
            run.font.size, run.font.bold, run.font.name, run.font.color.rgb = Pt(9.5), False, FONT, MID


def _scatter(slide, chart: dict, left, top, width, height):
    points = [p for p in chart.get("points") or [] if _is_number(p.get("x")) and _is_number(p.get("y"))]
    if not points:
        return None
    data = XyChartData()
    series = data.add_series(_clean(chart.get("title")) or "Точки")
    for p in points:
        series.add_data_point(p["x"], p["y"])
    frame = slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER, Inches(left), Inches(top), Inches(width), Inches(height), data)
    ch = frame.chart
    _style_axes(ch, min(_places([p["y"] for p in points]), 2), chart.get("unit") or "", chart, xy=True)
    ch.category_axis.tick_labels.number_format = _excel_format(min(_places([p["x"] for p in points]), 2))
    ch.category_axis.tick_labels.number_format_is_linked = False
    ch.has_legend = False
    ser = ch.plots[0].series[0]
    ser.format.line.fill.background()
    ser.marker.style = XL_MARKER_STYLE.CIRCLE
    ser.marker.size = 8
    ser.marker.format.fill.solid()
    ser.marker.format.fill.fore_color.rgb = RED
    ser.marker.format.line.color.rgb = WHITE
    if any(p.get("label") for p in points) and len(points) <= 20:
        for index, p in enumerate(points):
            if p.get("label"):
                label = ser.points[index].data_label
                label.has_text_frame = True
                label.text_frame.text = _cut(p["label"], 24)
                label.position = XL_LABEL_POSITION.RIGHT
                for para in label.text_frame.paragraphs:
                    for run in para.runs:
                        run.font.size, run.font.name, run.font.color.rgb = Pt(8.5), FONT, MID
    return frame


def _waterfall_bars(chart: dict) -> dict | None:
    """Раскладка водопада: основание и высота столбиков. Начало и итог посчитал код (ИИ-17)."""
    series = (chart.get("series") or [{}])[0]
    steps = [v if _is_number(v) else 0.0 for v in series.get("values") or []]
    xs = [str(x) for x in chart.get("x") or []][: len(steps)]
    if not steps:
        return None
    start, total = chart.get("start"), chart.get("total")
    places = min(_places(steps + [v for v in (start, total) if _is_number(v)]), 2)
    out = {"labels": [], "base": [], "bar": [], "colors": [], "texts": [], "places": places,
           "name": _clean(series.get("name")) or "Изменение"}

    def add(label, base, bar, color, text):
        for key, value in zip(("labels", "base", "bar", "colors", "texts"), (label, base, bar, color, text)):
            out[key].append(value)

    running = start if _is_number(start) else 0.0
    levels = [running]
    if _is_number(start):
        add("Начало", 0.0, start, MID, fmt_number(start, places))
    for x, v in zip(xs, steps):
        add(x, min(running, running + v), abs(v), RED if v < 0 else RGBColor(0xA6, 0xA6, 0xA6),
            fmt_number(v, places, signed=True))
        running += v
        levels.append(running)
    end = total if _is_number(total) else running
    add("Итог", 0.0, end, INK, fmt_number(end, places))
    levels.append(end)
    if min(out["base"] + levels) < 0:
        return None                      # водопад через ноль — обычными столбиками приращений
    # Изменения малы рядом с уровнем (выручка 12 млн, факторы по сотням тысяч): ось начинается не с нуля,
    # иначе столбики изменений не видны. На слайде об этом сказано.
    low, high = min(levels), max(levels)
    out["floor"] = None
    if _is_number(start) and low > 0 and high - low < 0.35 * high:
        span = max(high - low, 1e-9)
        step = 10 ** math.floor(math.log10(span))
        floor = math.floor((low - span * 0.6) / step) * step
        out["floor"] = floor if floor > 0 else None
    return out


def _waterfall_floor(chart: dict):
    if chart.get("type") != "waterfall":
        return None
    bars = _waterfall_bars(chart)
    return bars["floor"] if bars else None


def _waterfall(slide, chart: dict, left, top, width, height):
    bars = _waterfall_bars(chart)
    if bars is None:
        return None
    places = bars["places"]
    data = CategoryChartData(number_format=_excel_format(places))
    data.categories = bars["labels"]
    data.add_series("Основание", bars["base"])
    data.add_series(bars["name"], bars["bar"])
    frame = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, Inches(left), Inches(top), Inches(width),
                                   Inches(height), data)
    ch = frame.chart
    _style_axes(ch, places, chart.get("unit") or "", chart)
    if bars["floor"] is not None:
        ch.value_axis.minimum_scale = bars["floor"]
    ch.has_legend = False
    plot = ch.plots[0]
    plot.gap_width, plot.overlap = 50, 100
    hidden, shown = plot.series
    hidden.format.fill.background()
    hidden.format.line.fill.background()
    for index, color in enumerate(bars["colors"]):
        point = shown.points[index]
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = color
        if len(bars["labels"]) <= 16:
            label = point.data_label
            label.has_text_frame = True
            label.text_frame.text = bars["texts"][index]
            label.position = XL_LABEL_POSITION.INSIDE_END
            for para in label.text_frame.paragraphs:
                for run in para.runs:
                    run.font.size, run.font.name, run.font.bold = Pt(9), FONT, True
                    run.font.color.rgb = WHITE if color in (RED, MID, INK) else INK
    return frame


def _table_slide(deck: Deck, spec: dict, number: int):
    s = deck.slide()
    total_rows = len(spec["rows"])
    if spec.get("headline"):
        top = deck.frame(s, number, kicker="Главное", title=spec["headline"])
        deck.text(s, deck.L, top - 0.05, deck.CW, 0.3, spec["title"], size=13, bold=True, color=MID)
        top += 0.35
    else:
        top = deck.frame(s, number, kicker="Таблица", title=spec["title"])
    columns = spec["columns"][:TABLE_COLS]
    places = [_places([r[c] for r in spec["rows"] if c < len(r)]) for c in range(len(columns))]
    size = 11 if len(columns) <= 6 else 10
    limit = 60 if len(columns) <= 4 else (40 if len(columns) <= 6 else 28)
    cells = [[_cut(_cell(r[c] if c < len(r) else None, places[c]), limit) for c in range(len(columns))]
             for r in spec["rows"][:TABLE_ROWS]]
    numeric = [any(_is_number(r[c]) for r in spec["rows"][:TABLE_ROWS] if c < len(r)) for c in range(len(columns))]
    weights = [min(max(len(str(columns[c])) * 0.7, *(len(b[c]) for b in cells), 4), 34) for c in range(len(columns))]
    scale = deck.CW / sum(weights)
    widths = [w * scale for w in weights]
    line_h = size * 1.2 / 72
    head_h = max(0.4, 0.14 + line_h * max(_lines(str(c), w - 0.16, size, bold=True) for c, w in zip(columns, widths)))
    notes_h = 0.2 * 4
    room = deck.H - 0.62 - notes_h - top - head_h - 0.1
    body, heights, used = [], [], 0.0
    for values in cells:
        height = max(0.3, 0.1 + line_h * max(_lines(t, w - 0.16, size) for t, w in zip(values, widths)))
        if body and used + height > room:
            break
        body.append(values)
        heights.append(height)
        used += height
    rows = body
    shape = s.shapes.add_table(len(rows) + 1, len(columns), Inches(deck.L), Inches(top),
                               Inches(deck.CW), Inches(head_h + sum(heights)))
    table = shape.table
    _plain_table(shape)
    for c, width in enumerate(widths):
        table.columns[c].width = Inches(width)
    table.rows[0].height = Inches(head_h)
    for r, height in enumerate(heights, start=1):
        table.rows[r].height = Inches(height)
    for c, name in enumerate(columns):
        _fill_cell(table.cell(0, c), str(name), size=size, bold=True, color=WHITE, fill=RED,
                   align=PP_ALIGN.RIGHT if numeric[c] else PP_ALIGN.LEFT)
    for r, values in enumerate(body, start=1):
        for c, text in enumerate(values):
            _fill_cell(table.cell(r, c), text, size=size, color=INK, fill=PALE if r % 2 == 0 else WHITE,
                       align=PP_ALIGN.RIGHT if numeric[c] and text != "—" else PP_ALIGN.LEFT, bottom=BORDER)
    notes = [spec["source"]]
    if total_rows > len(rows) or spec["truncated"]:
        notes.append(f"Показаны первые {len(rows)} строк из {total_rows}"
                     + (" (результат обрезан лимитом строк)" if spec["truncated"] else "")
                     + " — полностью в выгрузке Excel.")
    if len(spec["columns"]) > TABLE_COLS:
        notes.append(f"Показаны {TABLE_COLS} столбцов из {len(spec['columns'])}.")
    if spec.get("note"):
        notes.append(spec["note"])
    y = deck.H - 0.62 - 0.2 * len(notes)
    deck.text(s, deck.L, y, deck.CW, 0.2 * len(notes), [[(n, {})] for n in notes], size=9.5, color=SOFT, spacing=1.05)


def _plain_table(shape):
    """Без темы таблицы: цвета задаются явно, иначе PowerPoint подмешает синие полосы темы."""
    tbl = shape._element.graphic.graphicData.tbl
    pr = tbl.tblPr
    for attr in ("firstRow", "bandRow"):
        pr.set(attr, "0")
    style = pr.find(qn("a:tableStyleId"))
    if style is None:
        style = pr.makeelement(qn("a:tableStyleId"), {})
        pr.append(style)
    style.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"      # «Без стиля, без сетки»


def _fill_cell(cell, text, *, size, color, fill, align, bold=False, bottom=None):
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    cell.margin_left = cell.margin_right = Inches(0.08)
    cell.margin_top = cell.margin_bottom = Inches(0.04)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame = cell.text_frame
    frame.word_wrap = True
    para = frame.paragraphs[0]
    para.alignment = align
    run = para.add_run()
    run.text = text
    run.font.name, run.font.size, run.font.bold, run.font.color.rgb = FONT, Pt(size), bold, color
    if bottom is not None:
        tc_pr = cell._tc.get_or_add_tcPr()
        line = tc_pr.makeelement(qn("a:lnB"), {"w": str(Pt(0.75)), "cap": "flat", "cmpd": "sng", "algn": "ctr"})
        solid = line.makeelement(qn("a:solidFill"), {})
        rgb = solid.makeelement(qn("a:srgbClr"), {"val": str(bottom)})
        solid.append(rgb)
        line.append(solid)
        # a:lnB идёт после a:lnL/a:lnR/a:lnT и до заливки ячейки — порядок схемы OOXML.
        fill_el = tc_pr.find(qn("a:solidFill"))
        if fill_el is not None:
            fill_el.addprevious(line)
        else:
            tc_pr.append(line)


def _points_slide(deck: Deck, spec: dict, number: int):
    s = deck.slide()
    suffix = f" ({spec['part']} из {spec['parts']})" if spec["parts"] > 1 else ""
    deck.frame(s, number, kicker="Выводы", title="Что показал анализ" + suffix)
    _draw_points(deck, s, spec["points"], BODY_TOP)
    legend = "Тип пункта ставит код по происхождению чисел: факт — из результата запроса; расчёт — по формуле из " \
             "чисел результата; гипотеза — требует проверки."
    if spec.get("more"):
        legend = spec["more"] + " " + legend
    deck.text(s, deck.L, deck.H - 0.86, deck.CW, 0.26, legend, size=9.5, color=SOFT)


def _recs_slide(deck: Deck, spec: dict, number: int):
    s = deck.slide()
    suffix = f" ({spec['part']} из {spec['parts']})" if spec["parts"] > 1 else ""
    deck.frame(s, number, kicker="Рекомендации", title="Что можно сделать" + suffix)
    y = BODY_TOP
    for rec in spec["recs"]:
        height = _rec_height(rec, deck.CW) - 0.2
        deck.rect(s, deck.L, y, 0.07, height, fill=RED)
        deck.tag(s, deck.L + 0.25, y + 0.02, "recommendation")
        action_h = _lines(rec["action"], deck.CW - REC_INDENT, 15, bold=True) * 15 * 1.25 / 72
        deck.text(s, deck.L + REC_INDENT, y, deck.CW - REC_INDENT, action_h, rec["action"], size=15, bold=True,
                  spacing=1.05)
        lines = []
        for key, label in (("basis", "Основание"), ("effect", "Ожидаемый эффект"), ("limits", "Когда не подходит")):
            if rec.get(key):
                lines.append([(f"{label}: ", {"bold": True, "color": INK}), (rec[key], {})])
        tail = "; ".join(t for t in (f"уверенность: {rec['confidence']}" if rec.get("confidence") else "",
                                     f"источник: {rec['source']}" if rec.get("source") else "") if t)
        if tail:
            lines.append([(tail[0].upper() + tail[1:], {"italic": True})])
        if lines:
            deck.text(s, deck.L + REC_INDENT, y + action_h + 0.06, deck.CW - REC_INDENT, height - action_h, lines,
                      size=11, color=MID, spacing=1.12)
        y += height + 0.2
    notes = [t for t in (spec.get("more"), spec.get("note")) if t]
    if notes:
        deck.text(s, deck.L, deck.H - 0.9, deck.CW, 0.3, " ".join(notes), size=10.5, color=MID)


def _limits_slide(deck: Deck, spec: dict, number: int):
    s = deck.slide()
    deck.frame(s, number, kicker="Ограничения", title="Ограничения и оговорки")
    y = BODY_TOP
    for item in spec["items"]:
        h = _lines(item, deck.CW - 0.4, 14) * 14 * 1.25 / 72
        if y + h > BODY_BOTTOM:
            break
        deck.rect(s, deck.L, y + 0.1, 0.09, 0.09, fill=RED)
        deck.text(s, deck.L + 0.35, y, deck.CW - 0.35, h, item, size=14, spacing=1.08)
        y += h + 0.18


def _sources_slide(deck: Deck, spec: dict, number: int):
    s = deck.slide()
    deck.frame(s, number, kicker="Источники", title="Источники и дата формирования")
    rows = spec["rows"]
    label_w, value_w = 2.9, deck.CW - 2.9
    heights = [max(0.3, 0.1 + 0.19 * _lines(v, value_w - 0.2, 10.5)) for _k, v in rows]
    shape = s.shapes.add_table(len(rows), 2, Inches(deck.L), Inches(BODY_TOP - 0.05), Inches(deck.CW),
                               Inches(sum(heights)))
    table = shape.table
    _plain_table(shape)
    table.columns[0].width, table.columns[1].width = Inches(label_w), Inches(value_w)
    for r, ((key, value), h) in enumerate(zip(rows, heights)):
        table.rows[r].height = Inches(h)
        _fill_cell(table.cell(r, 0), key, size=10.5, bold=True, color=INK, fill=PALE if r % 2 else WHITE,
                   align=PP_ALIGN.LEFT, bottom=BORDER)
        _fill_cell(table.cell(r, 1), value, size=10.5, color=INK if key != "Гриф" else RED,
                   fill=PALE if r % 2 else WHITE, align=PP_ALIGN.LEFT, bottom=BORDER)


def render(slides: list[dict], *, title: str = "", template: str | None = None) -> bytes:
    deck = Deck(template if template is not None else os.environ.get("AI_PPTX_TEMPLATE"))
    deck.total = len(slides)
    charts = [s for s in slides if s["kind"] == "chart"]
    chart_index = 0
    for number, spec in enumerate(slides, start=1):
        kind = spec["kind"]
        if kind == "title":
            _title_slide(deck, spec, number)
        elif kind == "headline":
            _headline_slide(deck, spec, number)
        elif kind == "chart":
            chart_index += 1
            _chart_slide(deck, spec, number, chart_index, len(charts))
        elif kind == "table":
            _table_slide(deck, spec, number)
        elif kind == "kpi":
            _kpi_slide(deck, spec, number)
        elif kind == "points":
            _points_slide(deck, spec, number)
        elif kind == "recs":
            _recs_slide(deck, spec, number)
        elif kind == "limits":
            _limits_slide(deck, spec, number)
        elif kind == "sources":
            _sources_slide(deck, spec, number)
    props = deck.prs.core_properties
    props.title = _cut(title, 200) or "Ответ ИИ-аналитика"
    props.author = "ИИ-аналитик (Инструмент АУП)"
    props.subject = GRIF
    props.keywords = "ИИ-аналитик; презентация; конфиденциально"
    props.comments = FOOTER
    props.last_modified_by = "ИИ-аналитик"
    return deck.save()


def build(message: dict, meta: dict, template: str | None = None) -> tuple[bytes, int]:
    """Презентация по ответу: файл и число слайдов (для журнала выгрузок)."""
    slides = outline(message, meta)
    return render(slides, title=message.get("question") or "", template=template), len(slides)


def file_name(message_id: int, question: str) -> str:
    words = re.findall(r"[A-Za-zА-Яа-яЁё0-9]+", question or "")
    slug = "_".join(words)[:48].rstrip("_") or "ответ"
    return f"Презентация_ИИ_{message_id}_{slug}.pptx"
