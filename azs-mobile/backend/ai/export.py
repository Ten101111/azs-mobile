"""ИИ-12. Выгрузка таблиц и данных графиков ответа ИИ в XLSX и CSV с паспортом.

Файл собирается на сервере из сохранённого ответа (`ai_messages.answer_json`), а
не из того, что нарисовано в браузере: в нём ровно те строки, что автор уже видел.
Выгрузить можно только свой ответ — владение проверяет `dialogs.message()`, как и
везде в диалогах (модель приватности ИБ-4: администратор чужую переписку не видит).

Паспорт — откуда цифры: вопрос, когда сформированы ответ и выгрузка, роль и область
данных, источник и таблицы, период, фильтры, тип данных (факт — строки витрины,
расчёт — посчитано по ним), версия каталога, номер ответа в журнале. SQL в паспорте —
только администратору и субадминистратору. Результат, обрезанный лимитом строк, —
«выгрузка неполная». Гриф «Конфиденциально» — в паспорте и колонтитулах.

Части ответа: `main` — основная таблица, `table-<id>` — дополнительные таблицы,
`chart-<id>` — данные графика.
"""
from __future__ import annotations

import csv
import hashlib
import io
import re
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MSK = timezone(timedelta(hours=3))
GRIF = "Конфиденциально"
FOOTER = "LUKOIL ОНПО · Инструмент АУП · Конфиденциально"
FORMATS = {"xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
           "csv": "text/csv; charset=utf-8"}
HEAD_FILL = PatternFill("solid", fgColor="E31E24")
HEAD_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(vertical="top", wrap_text=True)
DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
SQL_DATE_RE = re.compile(r"'(\d{4}-\d{2}-\d{2})'")
TASK_TITLES = {"lookup": "факт", "compare": "сравнение", "trend": "динамика", "diagnose": "диагностика причин",
               "anomaly": "поиск аномалий", "opportunity": "точки роста", "whatif": "сценарий", "other": "разбор"}
DEPTH_TITLES = {"fast": "Лёгкий", "analyze": "Средний", "deep": "Высокий"}


class NotExportable(ValueError):
    """В ответе нет такой таблицы или графика — выгружать нечего."""


# --- что выгружаем -------------------------------------------------------------

def _chart_table(chart: dict) -> tuple[list[str], list[list]]:
    """Данные графика таблицей: ровно те числа, по которым он нарисован."""
    kind = chart.get("type")
    unit = chart.get("unit") or ""
    if kind == "table":
        return list(chart.get("columns") or []), [list(r) for r in chart.get("rows") or []]
    if kind == "kpi":
        cards = chart.get("cards") or []
        with_delta = any("delta" in c for c in cards)
        columns = ["Показатель", "Значение"] + (["Изменение"] if with_delta else [])
        return columns, [[c.get("label"), c.get("value")] + ([c.get("delta")] if with_delta else []) for c in cards]
    if kind == "scatter":
        points = chart.get("points") or []
        labelled = any("label" in p for p in points)
        columns = (["Подпись"] if labelled else []) + [chart.get("xTitle") or "X", chart.get("yTitle") or "Y"]
        return columns, [([p.get("label")] if labelled else []) + [p.get("x"), p.get("y")] for p in points]
    series = chart.get("series") or []
    names = [s.get("name") or f"Ряд {i + 1}" for i, s in enumerate(series)]
    if unit:
        names = [n if unit in n else f"{n}, {unit}" for n in names]
    columns = [chart.get("xTitle") or "Ось X"] + names
    xs = chart.get("x") or []
    values = [list(s.get("values") or []) for s in series]
    rows = [[x] + [v[i] if i < len(v) else None for v in values] for i, x in enumerate(xs)]
    return columns, rows


def part_of(answer: dict, part: str) -> dict:
    """Таблица для выгрузки: название, столбцы, строки, источник (rN — витрина, pN — расчёт), усечение."""
    frame = answer.get("frame") or {}
    if part == "main":
        rows = answer.get("rows") or []
        if not rows:
            raise NotExportable("В ответе нет таблицы")
        source = str(frame.get("mainResult") or ("r1" if answer.get("depth", "fast") == "fast" else ""))
        step = _step(answer, source)
        return {"what": "основная таблица ответа",
                "title": (step or {}).get("purpose") or "Разбивка по строкам", "columns": list(answer.get("columns") or []),
                "rows": rows, "source": source, "truncated": bool(answer.get("truncated")), "note": "",
                "sql": answer.get("sql") or (step or {}).get("sql") or ""}
    kind, _, ident = part.partition("-")
    if kind == "table":
        table = next((t for t in answer.get("tables") or [] if str(t.get("id")) == ident), None)
        if table is None:
            raise NotExportable("В ответе нет такой таблицы")
        step = _step(answer, ident)
        return {"what": "таблица ответа", "title": table.get("title") or ident, "columns": list(table.get("columns") or []),
                "rows": list(table.get("rows") or []), "source": ident, "truncated": bool(table.get("truncated")),
                "note": "", "sql": (step or {}).get("sql") or ""}
    if kind == "chart":
        chart = next((c for c in answer.get("charts") or [] if str(c.get("id")) == ident), None)
        if chart is None:
            raise NotExportable("В ответе нет такого графика")
        columns, rows = _chart_table(chart)
        source = str(chart.get("source") or "")
        step = _step(answer, source)
        return {"what": "данные графика", "title": chart.get("title") or "График", "columns": columns, "rows": rows,
                "source": source,
                "truncated": "Показаны первые" in str(chart.get("note") or ""), "note": str(chart.get("note") or ""),
                "sql": (step or {}).get("sql") or ""}
    raise NotExportable("Неизвестная часть ответа")


def _step(answer: dict, source: str) -> dict | None:
    if not source:
        return None
    return next((s for s in answer.get("steps") or []
                 if str(s.get("resultId") or "").lower() == source.lower()), None)


def parts(answer: dict) -> list[str]:
    """Какие части ответа можно выгрузить — для кнопок и тестов."""
    out = ["main"] if answer.get("rows") else []
    out += [f"table-{t.get('id')}" for t in answer.get("tables") or [] if t.get("rows")]
    out += [f"chart-{c.get('id')}" for c in answer.get("charts") or []]
    return out


# --- паспорт -------------------------------------------------------------------

def _stamp(ts) -> str:
    return datetime.fromtimestamp(int(ts), MSK).strftime("%d.%m.%Y %H:%M МСК") if ts else "—"


def _tables_in(sql: str, dialect: str) -> list[str]:
    if not sql:
        return []
    try:
        import sqlglot
        from sqlglot import exp
        tree = sqlglot.parse_one(sql, read=dialect)
        names = {".".join(p for p in (t.db, t.name) if p) for t in tree.find_all(exp.Table)}
        ctes = {c.alias_or_name for c in tree.find_all(exp.CTE)}
        return sorted(n for n in names if n and n not in ctes)
    except Exception:  # noqa: BLE001 - без списка таблиц паспорт всё равно полезен
        return []


def _period(frame: dict, sql: str) -> str:
    if frame.get("period"):
        return str(frame["period"])
    days = sorted(set(SQL_DATE_RE.findall(sql or "")))
    if days:
        first, last = (datetime.strptime(d, "%Y-%m-%d").strftime("%d.%m.%Y") for d in (days[0], days[-1]))
        return f"по условию запроса: {first}" + (f" — {last}" if last != first else "")
    return "в вопросе не задан — по данным витрины на момент ответа"


def catalog_version(catalog) -> str:
    """Каталог витрины, по которому строился запрос: файл, дата изменения, отпечаток."""
    path = Path(getattr(catalog, "source", "") or "")
    if not path.is_file():
        return "встроенный каталог стенда"
    digest = hashlib.sha256(path.read_bytes()).hexdigest()[:8]
    changed = datetime.fromtimestamp(path.stat().st_mtime, MSK).strftime("%d.%m.%Y")
    return f"{path.name} от {changed}, отпечаток {digest}"


def passport(*, message: dict, table: dict, journal_row: dict | None, user, show_sql: bool,
             source_label: str, dialect: str, catalog_text: str, role_title: str, now: datetime | None = None) -> list[tuple[str, str]]:
    answer = message["answer"]
    frame = answer.get("frame") or {}
    journal_row = journal_row or {}
    server_sql = table.get("sql") or journal_row.get("sql_final") or ""
    rows = len(table["rows"])
    fact = not str(table.get("source") or "").lower().startswith("p")
    depth = answer.get("depth") or "fast"
    params = [f"уровень «{DEPTH_TITLES.get(depth, depth)}»",
              f"задача — {TASK_TITLES.get(answer.get('taskType') or 'lookup', answer.get('taskType') or 'факт')}"]
    params += [str(f) for f in frame.get("filters") or []]
    binding = journal_row.get("binding") or ""
    actor = journal_row.get("actor") or ""
    role = role_title + (f" ({binding})" if binding else "")
    if "(от имени:" in actor:
        role += " — вопрос задан администратором от имени этой роли"
    tables = _tables_in(server_sql, dialect)
    out = [
        ("Гриф", f"{GRIF}. Не пересылайте за пределы компании."),
        ("Вопрос", message.get("question") or answer.get("question") or ""),
        ("Что выгружено", f"{table['what']}: «{table['title']}»"),
        ("Ответ сформирован", _stamp(message.get("createdAt"))),
        ("Выгрузка сформирована", (now or datetime.now(MSK)).astimezone(MSK).strftime("%d.%m.%Y %H:%M МСК")
         + (f", {getattr(user, 'email', '')}" if getattr(user, "email", "") else "")),
        ("Роль и область данных", f"{role}; область: {answer.get('scopeLabel') or journal_row.get('scope_label') or '—'} "
                                  "— подставлена системой, а не выбрана моделью"),
        ("Источник", source_label + (f": {', '.join(tables)}" if tables else "")),
        ("Период", _period(frame, server_sql)),
        ("Показатели", ", ".join(map(str, frame.get("metrics") or [])) or "—"),
        ("Фильтры и параметры", "; ".join(params)),
        ("Тип данных", "факт — строки, прочитанные из витрины" if fact
         else "расчёт — посчитано по строкам витрины (Python в песочнице)"),
        ("Строк в выгрузке", f"{rows}" + ("; выгрузка неполная: результат обрезан лимитом строк, в витрине строк больше"
                                          if table.get("truncated") else "")),
        ("Версия каталога", catalog_text),
        ("Номер ответа", f"сообщение {message.get('id')}" + (f", запись журнала {message.get('journalId')}"
                                                            if message.get("journalId") else "")),
    ]
    if table.get("note"):
        out.append(("Примечание графика", table["note"]))
    if show_sql and (table.get("sql") or journal_row.get("sql_final")):
        out.append(("SQL (администратор и субадминистратор)", table.get("sql") or journal_row.get("sql_final")))
    return out


# --- файлы -----------------------------------------------------------------------

def _decimals(values: list) -> int | None:
    """Знаков после запятой в столбце: None — столбец не числовой."""
    numbers = [v for v in values if isinstance(v, (int, float)) and not isinstance(v, bool)]
    if not numbers:
        return None
    places = 0
    for v in numbers:
        if isinstance(v, float) and not v.is_integer():
            text = f"{v:.6f}".rstrip("0")
            places = max(places, len(text.split(".")[1]) if "." in text else 0)
    return min(places, 3)


def _cell(value):
    if isinstance(value, str) and DATE_RE.match(value):
        try:
            return date.fromisoformat(value)
        except ValueError:
            return value
    return value


def build_xlsx(table: dict, passport_rows: list[tuple[str, str]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Данные"
    ws.append(table["columns"])
    for cell in ws[1]:
        cell.font, cell.fill = HEAD_FONT, HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in table["rows"]:
        ws.append([_cell(v) for v in row])
    for col, name in enumerate(table["columns"], start=1):
        values = [row[col - 1] for row in table["rows"] if col - 1 < len(row)]
        places = _decimals(values)
        number_format = None if places is None else ("#,##0" if places == 0 else "#,##0." + "0" * places)
        for r in range(2, ws.max_row + 1):
            cell = ws.cell(r, col)
            if isinstance(cell.value, date):
                cell.number_format = "DD.MM.YYYY"
            elif number_format and isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                cell.number_format = number_format
        width = max([len(str(name))] + [len(str(v)) for v in values[:200]] + [8])
        ws.column_dimensions[get_column_letter(col)].width = min(48, width + 2)
    ws.freeze_panes = "A2"

    info = wb.create_sheet("Паспорт")
    info.append(["Поле", "Значение"])
    for cell in info[1]:
        cell.font, cell.fill = HEAD_FONT, HEAD_FILL
    for key, value in passport_rows:
        info.append([key, value])
        info.cell(info.max_row, 2).alignment = WRAP
        info.cell(info.max_row, 1).alignment = WRAP
    info.column_dimensions["A"].width = 30
    info.column_dimensions["B"].width = 110
    for sheet in (ws, info):
        sheet.oddHeader.right.text = GRIF
        sheet.oddFooter.left.text = FOOTER
    wb.properties.title = table["title"]
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _csv_value(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "да" if value else "нет"
    if isinstance(value, float):
        text = f"{value:.6f}".rstrip("0").rstrip(".") if not value.is_integer() else str(int(value))
        return text.replace(".", ",")
    if isinstance(value, str) and DATE_RE.match(value):
        return datetime.strptime(value, "%Y-%m-%d").strftime("%d.%m.%Y")
    return str(value)


def build_csv(table: dict, passport_rows: list[tuple[str, str]]) -> bytes:
    """CSV для русского Excel: UTF-8 с BOM, разделитель «;», десятичная запятая. Паспорт — под таблицей."""
    out = io.StringIO()
    writer = csv.writer(out, delimiter=";", lineterminator="\r\n")
    writer.writerow([f"{GRIF} — выгрузка ответа ИИ-аналитика"])
    writer.writerow([])
    writer.writerow(table["columns"])
    for row in table["rows"]:
        writer.writerow([_csv_value(v) for v in row])
    writer.writerow([])
    writer.writerow(["Паспорт выгрузки"])
    for key, value in passport_rows:
        writer.writerow([key, " ".join(str(value).split())])
    return out.getvalue().encode("utf-8-sig")


def file_name(message_id: int, title: str, fmt: str) -> str:
    words = re.findall(r"[A-Za-zА-Яа-яЁё0-9]+", title or "")
    slug = "_".join(words)[:48].rstrip("_") or "таблица"
    return f"Ответ_ИИ_{message_id}_{slug}.{fmt}"
