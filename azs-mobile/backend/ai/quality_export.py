"""Выгрузка раздела «Качество ответов» в Excel (БТ-КК9).

Выгружается то же, что видно на экране: сводка, лента оценок и срез по
версиям. Переписки в файле нет — выгрузка не должна давать больше, чем
интерфейс, иначе модель приватности ИБ-4 обходится через кнопку «Скачать».
"""
from __future__ import annotations

import io
import time

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import quality

HEAD_FILL = PatternFill("solid", fgColor="17191D")
HEAD_FONT = Font(color="F5F5F3", bold=True, size=10)
WRAP = Alignment(vertical="top", wrap_text=True)


def _date(value) -> str:
    if not value:
        return ""
    return time.strftime("%d.%m.%Y %H:%M", time.localtime(int(value)))


def _sheet(wb: Workbook, title: str, head: list[str], widths: list[int]):
    ws = wb.create_sheet(title)
    ws.append(head)
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    return ws


def build(days: int = 30) -> bytes:
    summary = quality.summary(days)
    rows = quality.entries(days=days, limit=1000)
    versions = quality.versions()

    wb = Workbook()
    wb.remove(wb.active)

    ws = _sheet(wb, "Сводка", ["Показатель", "Значение"], [46, 22])
    share = lambda part, whole: f"{round(100.0 * part / whole, 1)} %" if whole else "—"
    for name, value in (
        ("Период, дней", summary["days"]),
        ("Задано вопросов", summary["asked"]),
        ("Отвечено", summary["answered"]),
        ("Отказов (граница области данных)", summary["refused"]),
        ("Технических ошибок", summary["failed"]),
        ("Доля отказов", share(summary["refused"], summary["asked"])),
        ("Доля технических ошибок", share(summary["failed"], summary["asked"])),
        ("Оценено ответов", summary["rated"]),
        ("Доля оценённых", share(summary["rated"], summary["answered"])),
        ("Средняя оценка", summary["average"] if summary["average"] is not None else "—"),
        ("Среднее время ответа, с", round(summary["avgMs"] / 1000, 1)),
        ("Наибольшее время ответа, с", round(summary["maxMs"] / 1000, 1)),
        ("В разборе и не разобрано", summary["openReview"]),
        (f"Просрочено (реакция дольше {summary['firstResponseDays']} дней)", summary["overdue"]),
    ):
        ws.append([name, value])
    ws.append([])
    ws.append(["Распределение оценок", ""])
    for star in range(1, 6):
        ws.append([f"{star} звёзд", summary["spread"][str(star)]])

    ws = _sheet(
        wb, "Оценки",
        ["Дата", "Оценка", "Роль", "Область данных", "Вопрос", "Комментарий",
         "Исход", "Правило", "Модель", "Версия инструкции", "Время, с",
         "Статус разбора", "Ответственный", "Результат разбора", "В эталонном наборе"],
        [17, 8, 10, 26, 52, 46, 12, 18, 18, 16, 10, 14, 20, 40, 18],
    )
    verdicts = {"ok": "ответ", "rejected": "отказ",
                "execution_error": "ошибка", "model_unavailable": "ошибка"}
    for row in rows:
        ws.append([
            _date(row["created_at"]), row["rating"], row["role"] or "",
            row["scope_label"] or "", row["question"], row["comment"] or "",
            verdicts.get(row["verdict"], row["verdict"] or ""), row["rule"] or "",
            row["model"] or "", row["prompt_version"] or "",
            round((row["total_ms"] or 0) / 1000, 1),
            row["statusTitle"], row["owner"] or "", row["note"] or "",
            "да" if row["in_golden"] else "",
        ])
    for line in ws.iter_rows(min_row=2):
        for cell in line:
            cell.alignment = WRAP

    ws = _sheet(
        wb, "Версии",
        ["Версия инструкции", "Модель", "Задано", "Отвечено", "Оценено",
         "Средняя оценка", "Низких оценок", "С", "По"],
        [20, 20, 12, 12, 12, 16, 16, 17, 17],
    )
    for item in versions:
        ws.append([
            item["version"], item["model"] or "", item["asked"], item["answered"],
            item["rated"], item["average"] if item["average"] is not None else "—",
            item["low"], _date(item["sinceAt"]), _date(item["untilAt"]),
        ])

    ws = _sheet(wb, "О выгрузке", ["", ""], [110, 2])
    for line in (
        "Оценки ответов ИИ — обратная связь по продукту.",
        "Они не используются для оценки работы сотрудника и не влияют на премирование (БТ-КК10).",
        "",
        "Как только оценка начнёт попадать в разговор о работе человека, низкие оценки исчезнут:",
        "ставить их станет невыгодно. Продукт при этом не станет лучше — он просто перестанет",
        "получать сигнал о том, где врёт.",
        "",
        "Переписки в выгрузке нет: доступны метаданные, оценка, комментарий и текст запроса",
        "к витрине. Полное содержание чужого диалога раскрывается только по жалобе его автора",
        "(модель приватности ИБ-4, БТ-КК8).",
        "",
        f"Выгружено: {time.strftime('%d.%m.%Y %H:%M')}. Период: {days} дней.",
    ):
        ws.append([line, ""])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
