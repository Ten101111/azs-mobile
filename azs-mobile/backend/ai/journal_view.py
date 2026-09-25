"""«Журнал ИИ» в админке (25.09.2026): все обращения к ИИ-аналитику с фильтрами.

Читает журнал ai_queries и тематики из backend/ai/topics.py. Фильтры: период,
роль, уровень, исход, текст вопроса и пять тематических фильтров. Внутри
одного тематического фильтра значения объединяются через «или», между
фильтрами — через «и». Счётчик у значения фильтра считается при всех прочих
выбранных фильтрах, кроме своего, — видно, сколько найдётся, если его выбрать.

Только администратору (api.py, require_admin): здесь вопросы людей и SQL.
"""
from __future__ import annotations

import io
import json
import sqlite3
import time
from typing import Any

from . import topics

DAY = 86400
DEPTH_TITLES = {"fast": "Лёгкий", "analyze": "Средний", "deep": "Высокий", "auto": "Авто"}
OUTCOMES = {
    "ok": ("Ответ", ("ok",)),
    "refused": ("Отказ", ("rejected",)),
    "failed": ("Ошибка", ("execution_error", "model_unavailable")),
    "nodata": ("Нет данных", ("agent_no_data",)),
}
KNOWN_VERDICTS = tuple(v for _title, codes in OUTCOMES.values() for v in codes)
STEP_TITLES = {"schema": "Справка по схеме", "sql": "Запрос к витрине", "python": "Расчёт",
               "chart": "График", "write": "Итоговый текст"}
MAX_LIMIT = 200


def outcome_of(verdict: str | None, rule: str | None = None) -> str:
    if rule == "cancelled":
        return "cancelled"
    for code, (_title, verdicts) in OUTCOMES.items():
        if verdict in verdicts:
            return code
    return "other"


def outcome_title(code: str) -> str:
    return {"cancelled": "Остановлен", "other": "Другое"}.get(code) or OUTCOMES[code][0]


def _conn() -> sqlite3.Connection:
    conn = topics.connect()
    conn.create_function("casefold", 1, lambda v: (v or "").casefold())
    return conn


def _where(days: int, role: str, depth: str, outcome: str, text: str,
           chosen: dict[str, list[str]], skip: str | None = None, now: int | None = None) -> tuple[str, list]:
    where, args = ["1 = 1"], []
    if days and days > 0:
        where.append("q.created_at >= ?")
        args.append(int(now if now is not None else time.time()) - int(days) * DAY)
    if role:
        where.append("q.role = ?")
        args.append(role)
    if depth:
        where.append("q.depth = ?")
        args.append(depth)
    if outcome == "cancelled":
        where.append("q.rule = 'cancelled'")
    elif outcome == "other":
        where.append(f"COALESCE(q.verdict, '') NOT IN ({','.join('?' * len(KNOWN_VERDICTS))})")
        args.extend(KNOWN_VERDICTS)
    elif outcome in OUTCOMES:
        codes = OUTCOMES[outcome][1]
        where.append(f"q.verdict IN ({','.join('?' * len(codes))}) AND COALESCE(q.rule, '') != 'cancelled'")
        args.extend(codes)
    if text.strip():
        where.append("casefold(q.question) LIKE ?")
        args.append(f"%{text.strip().casefold()}%")
    for facet, values in chosen.items():
        if facet == skip or not values:
            continue
        where.append(f"q.id IN (SELECT query_id FROM ai_query_topics WHERE facet = ? "
                     f"AND value IN ({','.join('?' * len(values))}))")
        args.extend([facet, *values])
    return " AND ".join(where), args


def parse_topics(items: list[str] | None) -> dict[str, list[str]]:
    """«facet:значение» из адреса → {facet: [значения]}; незнакомые фильтры отбрасываются."""
    out: dict[str, list[str]] = {}
    for item in items or []:
        facet, _, value = (item or "").partition(":")
        if facet in topics.FACET_TITLES and value.strip():
            out.setdefault(facet, [])
            if value.strip() not in out[facet]:
                out[facet].append(value.strip())
    return out


def _median(values: list[int]) -> int:
    if not values:
        return 0
    values = sorted(values)
    middle = len(values) // 2
    return values[middle] if len(values) % 2 else (values[middle - 1] + values[middle]) // 2


def search(days: int = 30, role: str = "", depth: str = "", outcome: str = "", text: str = "",
           chosen: dict[str, list[str]] | None = None, limit: int = 50, offset: int = 0,
           now: int | None = None) -> dict:
    chosen = chosen or {}
    now = int(now if now is not None else time.time())
    limit = max(1, min(int(limit), MAX_LIMIT))
    offset = max(0, int(offset))
    conn = _conn()
    try:
        sql, args = _where(days, role, depth, outcome, text, chosen, now=now)
        stats = conn.execute(f"SELECT q.verdict, q.rule, q.total_ms FROM ai_queries q WHERE {sql}", args).fetchall()
        rows = conn.execute(
            "SELECT q.id, q.created_at, q.actor, q.role, q.question, q.depth, q.depth_requested, q.verdict, q.rule, "
            "q.total_ms, q.model_ms, q.sql_ms, q.model_calls, q.tool_calls, q.stop_reason, q.row_count, q.scope_label "
            f"FROM ai_queries q WHERE {sql} ORDER BY q.id DESC LIMIT ? OFFSET ?", [*args, limit, offset]).fetchall()
        tags = topics.query_topics(conn, [r["id"] for r in rows])
        facets = []
        fresh_since = now - topics.NEW_DAYS * DAY
        for facet, title in topics.FACETS:
            f_sql, f_args = _where(days, role, depth, outcome, text, chosen, skip=facet, now=now)
            values = conn.execute(
                "SELECT t.value, COUNT(DISTINCT t.query_id) AS n, MIN(r.first_seen) AS first_seen, "
                "MAX(r.origin) AS origin FROM ai_query_topics t JOIN ai_queries q ON q.id = t.query_id "
                "LEFT JOIN ai_topics r ON r.facet = t.facet AND r.value = t.value "
                f"WHERE t.facet = ? AND {f_sql} GROUP BY t.value ORDER BY n DESC, t.value",
                [facet, *f_args]).fetchall()
            items = [{"value": v["value"], "count": int(v["n"]), "origin": v["origin"] or "rules",
                      "isNew": bool(v["first_seen"] and int(v["first_seen"]) >= fresh_since)} for v in values]
            present = {item["value"] for item in items}
            items += [{"value": value, "count": 0, "origin": "rules", "isNew": False}
                      for value in chosen.get(facet, []) if value not in present]
            facets.append({"code": facet, "title": title, "values": items})
        roles_present = [r[0] for r in conn.execute("SELECT DISTINCT role FROM ai_queries WHERE role IS NOT NULL "
                                                    "ORDER BY role")]
        pending = int(conn.execute("SELECT COUNT(*) FROM ai_topic_state WHERE model_state = 'pending'").fetchone()[0])
    finally:
        conn.close()
    outcomes = [outcome_of(v, r) for v, r, _t in stats]
    return {
        "total": len(stats),
        "summary": {
            "total": len(stats),
            "ok": outcomes.count("ok"),
            "refused": outcomes.count("refused"),
            "failed": outcomes.count("failed"),
            "medianMs": _median([int(t) for _v, _r, t in stats if t]),
        },
        "entries": [_entry(dict(r), tags.get(int(r["id"]), {})) for r in rows],
        "facets": facets,
        "roles": [{"code": code, "title": _role_title(code)} for code in roles_present],
        "depths": [{"code": c, "title": t} for c, t in DEPTH_TITLES.items() if c != "auto"],
        "outcomes": [{"code": c, "title": t} for c, (t, _v) in OUTCOMES.items()]
                    + [{"code": "cancelled", "title": "Остановлен"}, {"code": "other", "title": "Другое"}],
        "pendingThemes": pending,
        "modelThemes": topics.MODEL_ENABLED,
        "offset": offset,
        "limit": limit,
    }


def _role_title(code: str | None) -> str:
    try:
        from .. import roles

        spec = roles.ROLES.get(code or "")
        return spec.title if spec else (code or "—")
    except Exception:  # noqa: BLE001
        return code or "—"


def _entry(row: dict, tags: dict[str, list[str]]) -> dict:
    outcome = outcome_of(row.get("verdict"), row.get("rule"))
    return {
        "id": row["id"],
        "at": row["created_at"],
        "actor": row.get("actor") or "",
        "role": row.get("role") or "",
        "roleTitle": _role_title(row.get("role")),
        "question": row.get("question") or "",
        "depth": row.get("depth") or "",
        "depthTitle": DEPTH_TITLES.get(row.get("depth") or "", row.get("depth") or "—"),
        "depthRequested": row.get("depth_requested") or "",
        "outcome": outcome,
        "outcomeTitle": outcome_title(outcome),
        "verdict": row.get("verdict") or "",
        "totalMs": row.get("total_ms") or ((row.get("model_ms") or 0) + (row.get("sql_ms") or 0)) or 0,
        "modelCalls": row.get("model_calls"),
        "toolCalls": row.get("tool_calls"),
        "stopReason": row.get("stop_reason") or "",
        "rows": row.get("row_count"),
        "scope": row.get("scope_label") or "",
        "topics": tags,
    }


def detail(query_id: int) -> dict | None:
    conn = _conn()
    try:
        row = conn.execute("SELECT * FROM ai_queries WHERE id = ?", (int(query_id),)).fetchone()
        if row is None:
            return None
        tags = topics.query_topics(conn, [int(query_id)]).get(int(query_id), {})
        state = conn.execute("SELECT model_state FROM ai_topic_state WHERE query_id = ?", (int(query_id),)).fetchone()
        origins = {(r["facet"], r["value"]): r["origin"] for r in conn.execute(
            "SELECT facet, value, origin FROM ai_query_topics WHERE query_id = ?", (int(query_id),))}
    finally:
        conn.close()
    row = dict(row)
    entry = _entry(row, tags)
    try:
        trace = json.loads(row.get("trace_json") or "null") or {}
    except ValueError:
        trace = {}
    plan = trace.get("plan") if isinstance(trace, dict) else None
    plan = plan if isinstance(plan, dict) else {}
    steps = []
    for step in (trace.get("steps") if isinstance(trace, dict) else None) or []:
        if not isinstance(step, dict):
            continue
        steps.append({"kind": step.get("kind") or "", "kindTitle": STEP_TITLES.get(step.get("kind") or "", step.get("kind") or ""),
                      "label": step.get("label") or step.get("purpose") or "", "ms": step.get("ms") or 0,
                      "ok": step.get("ok", True), "error": step.get("error") or "", "rows": step.get("rows")})
    try:
        model_trace = json.loads(row.get("model_trace") or "null") or {}
    except ValueError:
        model_trace = {}
    metric_titles = [t for t in (topics.metric_title(k) for k in plan.get("metrics") or [] if isinstance(k, str)) if t]
    entry.update({
        "standalone": plan.get("standaloneQuestion") or "",
        "planPeriod": plan.get("period") or "",
        "planMetrics": [title for _k, title in metric_titles],
        "steps": steps,
        "sql": row.get("sql_final") or "",
        "rule": row.get("rule") or "",
        "message": row.get("message") or "",
        "model": row.get("model") or "",
        "promptVersion": row.get("prompt_version") or "",
        "modelMs": row.get("model_ms") or 0,
        "sqlMs": row.get("sql_ms") or 0,
        "promptMs": row.get("model_prompt_ms"),
        "evalMs": row.get("model_eval_ms"),
        "loadMs": row.get("model_load_ms"),
        "tokensIn": row.get("tokens_in"),
        "tokensOut": row.get("tokens_out"),
        "modelSteps": (model_trace.get("calls") if isinstance(model_trace, dict) else None) or [],
        # ИИ-07 / ИИ-11: файлы вопроса (без содержимого) и память папки.
        "files": json.loads(row.get("files_json") or "[]"),
        "memoryFolder": row.get("memory_folder") or "",
        "themeBy": "model" if origins.get(("theme", (tags.get("theme") or [""])[0])) == "model" else "rules",
        "themeState": state["model_state"] if state else "",
    })
    return entry


def export(days: int = 30, role: str = "", depth: str = "", outcome: str = "", text: str = "",
           chosen: dict[str, list[str]] | None = None) -> bytes:
    """Выгрузка того, что видно в списке: без SQL и шагов — они только на экране по клику."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    data = search(days, role, depth, outcome, text, chosen, limit=MAX_LIMIT)
    rows = list(data["entries"])
    offset = MAX_LIMIT
    while len(rows) < data["total"] and offset < 20_000:
        rows += search(days, role, depth, outcome, text, chosen, limit=MAX_LIMIT, offset=offset)["entries"]
        offset += MAX_LIMIT
    wb = Workbook()
    ws = wb.active
    ws.title = "Журнал ИИ"
    head = ["Дата и время", "Пользователь", "Роль", "Вопрос", "Уровень", "Исход", "Время ответа, с",
            "Вызовов модели", "Шагов", "Остановка", *[title for _c, title in topics.FACETS]]
    widths = [17, 28, 18, 70, 10, 12, 12, 10, 8, 18, 26, 30, 22, 18, 20]
    ws.append(head)
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="17191D")
        cell.font = Font(color="F5F5F3", bold=True, size=10)
    ws.freeze_panes = "A2"
    for item in rows:
        ws.append([
            time.strftime("%d.%m.%Y %H:%M", time.localtime(int(item["at"]))), item["actor"], item["roleTitle"],
            item["question"], item["depthTitle"], item["outcomeTitle"],
            round((item["totalMs"] or 0) / 1000, 1), item["modelCalls"], item["toolCalls"], item["stopReason"],
            *["; ".join(item["topics"].get(code, [])) for code, _t in topics.FACETS],
        ])
    for line in ws.iter_rows(min_row=2):
        for cell in line:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.append([])
    ws.append(["Конфиденциально. Вопросы пользователей к ИИ-аналитику — только для администратора."])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
