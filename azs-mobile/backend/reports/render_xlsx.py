"""СП-05. XLSX справки: листы по разделам и «Паспорт», числа — числами с форматом."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import fmt
from . import narrative
from .render_pdf import (SERVICE_NOTE, passport_lines, plan_caption, plan_months, service_month_line, service_title,
                         version_label)

RED = "E31E24"
HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor=RED)
TOTAL_FILL = PatternFill("solid", fgColor="F4F5F7")
INT = "#,##0"
DEC1 = "#,##0.0"
DEC2 = "#,##0.00"
DELTA = '+0.0;-0.0;0.0'


def _sheet(wb, title: str, header: list[str], rows: list[list], formats: list[str | None], widths: list[int],
           total_last: bool = False):
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font, cell.fill = HEAD_FONT, HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in rows:
        ws.append(row)
    for col, number_format in enumerate(formats, start=1):
        if not number_format:
            continue
        for r in range(2, ws.max_row + 1):
            ws.cell(r, col).number_format = number_format
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    if total_last and rows:
        for cell in ws[ws.max_row]:
            cell.fill = TOTAL_FILL
            cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    return ws


def _fmt(decimals: int) -> str:
    return {0: INT, 1: DEC1}.get(decimals, DEC2)


def render(model: dict, path: Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)

    # Цифры недели: изменения в % и п. п. — отдельными столбцами с подписью единицы.
    ws = wb.create_sheet("Цифры недели")
    ws.append(["Показатель", "Единица", "Неделя", "Прошлая неделя", "Δ н/н", "Единица Δ",
               "Та же неделя прошлого года", "Δ г/г"])
    for cell in ws[1]:
        cell.font, cell.fill = HEAD_FONT, HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for m in model["metrics"]:
        ws.append([m["title"], fmt.unit(m["unit"]), m["value"], m["prev"], m["deltaPrev"],
                   "п. п." if m["isShare"] else "%", m["lastYear"], m["deltaYear"]])
        r = ws.max_row
        for col in (3, 4, 7):
            ws.cell(r, col).number_format = _fmt(m["decimals"])
        for col in (5, 8):
            ws.cell(r, col).number_format = DELTA
    for col, width in enumerate([28, 10, 16, 16, 10, 10, 18, 10], start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"
    ws.append([])
    ws.append([f"{model['title']} · {model['scopeLabel']} · неделя {model['week']['label']} · "
               f"выпуск {model['week']['iso']}, {version_label(model)}"])
    for sentence in narrative.headline(model):
        ws.append([sentence])
    notes = narrative.attention(model)
    if notes:
        ws.append([])
        ws.append(["На что обратить внимание"])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        for line in notes:
            ws.append([f"• {line}"])
    recs = narrative.recommendations(model)
    if recs:
        ws.append([])
        ws.append(["Рекомендации (справочно)"])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        for rec in recs:
            ws.append([f"• {rec['action']} Основание: {rec['basis']} Ожидаемый эффект: {rec['effect']} "
                       f"Ограничения: {rec['limits']}"])
        ws.append([narrative.RECOMMENDATION_NOTE])

    months = plan_months(model)
    if months:
        ws = wb.create_sheet("План месяца")
        ws.append(["Месяц", "Показатель", "Единица", "Факт с начала месяца", "План месяца", "План на дату",
                   "Выполнение плана месяца, %", "Выполнение на дату, %", "Отставание, п. п.", "Нужно в день",
                   "Сейчас в день", "Статус"])
        for cell in ws[1]:
            cell.font, cell.fill = HEAD_FONT, HEAD_FILL
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for month in months:
            for r in month["rows"]:
                ws.append([plan_caption(month), r["title"], fmt.unit(r["unit"]), r["fact"], r["planMonth"],
                           r["planToDate"], r["pctMonth"], r["pctToDate"], r["gapPp"], r["needPerDay"],
                           r["pacePerDay"], r.get("status")])
                row = ws.max_row
                for col in (4, 5, 6, 10, 11):
                    ws.cell(row, col).number_format = _fmt(r["decimals"])
                for col in (7, 8):
                    ws.cell(row, col).number_format = DEC1
                ws.cell(row, 9).number_format = DELTA
            ws.append([])
            ws.append([f"{plan_caption(month)} — ОНПО", "ОНПО", "Топливо, % плана", "Выручка НТУ, % плана",
                       "НТУ, отставание п. п.", "ВД НТУ, % плана", "Статус"])
            ws.cell(ws.max_row, 1).font = Font(bold=True)
            for o in month["onpo"]:
                ws.append(["", o["name"], o.get("fuelPct"), o.get("ntuPct"), o.get("ntuGap"), o.get("vdPct"),
                           o.get("status")])
            ws.append([])
        for col, width in enumerate([34, 22, 10, 18, 18, 16, 16, 16, 14, 16, 16, 12], start=1):
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = "A2"

    service = model.get("service")
    if service:
        ws = wb.create_sheet("Сервис")
        ws.append(["Показатель", "Неделя", "Прошлая неделя", "Δ н/н", "Прошлый год", "Δ г/г", "Изменение"])
        for cell in ws[1]:
            cell.font, cell.fill = HEAD_FONT, HEAD_FILL
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        kinds = {"abs": "разность", "pct": "%", "pp": "п. п."}
        for m in service["metrics"]:
            ws.append([service_title(m), m["value"], m["prev"], m["deltaPrev"], m["lastYear"], m["deltaYear"],
                       kinds.get(m["deltaKind"], "")])
            for col in (2, 3, 4, 5, 6):
                ws.cell(ws.max_row, col).number_format = _fmt(m["decimals"]) if m["decimals"] <= 2 else "0.000"
        for month in service.get("months", []):
            ws.append([])
            ws.append([service_month_line(month)])
        if service.get("categories"):
            ws.append([])
            ws.append(["Негатив по категориям", "Неделя", "Прошлая неделя"])
            ws.cell(ws.max_row, 1).font = Font(bold=True)
            for c in service["categories"]:
                ws.append([c["title"], c["week"], c["prev"]])
        if service.get("onpo"):
            ws.append([])
            ws.append(["ОНПО", "АЗС с оценками", "Средняя оценка", "Δ н/н", "Негативных", "Жалоб ЕГЛ",
                       "Качество сервиса, на 100 тыс. чеков"])
            ws.cell(ws.max_row, 1).font = Font(bold=True)
            for o in service["onpo"]:
                ws.append([o["name"], o["stations"], o["avg"], o.get("avgDeltaPrev"), o["negative"], o.get("complaints"),
                           o.get("quality")])
                ws.cell(ws.max_row, 3).number_format = "0.000"
                ws.cell(ws.max_row, 4).number_format = "+0.000;-0.000;0.000"
                ws.cell(ws.max_row, 7).number_format = DEC2
        ws.append([])
        ws.append([SERVICE_NOTE])
        for col, width in enumerate([40, 16, 16, 14, 16, 14, 14], start=1):
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = "A2"

    unit = fmt.unit(model["fuelUnit"])
    rows = [[r["name"], r.get("stations"), r.get("fuel"), r.get("fuelDeltaPrev"), r.get("fuelDeltaYear"),
             r.get("ntu"), r.get("ntuDeltaPrev"), r.get("ntuDeltaYear"), r.get("conversion")]
            for r in model["onpo"] + [dict(model["onpoTotal"], name="Итого по сети")]]
    _sheet(wb, "ОНПО", ["ОНПО", "Объектов", f"Топливо, {unit}", "Δ н/н, %", "Δ г/г, %", "Выручка НТУ, ₽",
                        "Δ н/н, %", "Δ г/г, %", "Конверсия НТУ, %"],
           rows, [None, INT, INT, DELTA, DELTA, INT, DELTA, DELTA, DEC1], [22, 10, 16, 10, 10, 18, 10, 10, 14],
           total_last=True)

    att = model["attention"]
    rows = [["Падение топлива", r["label"], r["region"], r["onpo"], r["fuel"], r["fuelPrev"], r["delta"], None]
            for r in att["drops"]]
    rows += [["Без продаж", r["label"], r["region"], r["onpo"], None, None, None, r["idleDays"]] for r in att["noSales"]]
    rows += [["Лидер роста", r["label"], r["region"], r["onpo"], r["fuel"], r["fuelPrev"], r["delta"], None]
             for r in att["leaders"]]
    rows = [r + [None, None, None] for r in rows]
    rows += [["Негатив в приложении", r["label"], r["region"], r["onpo"], None, None, None, None, r["negative"],
              r["avg"], r["category"]] for r in att.get("negative", [])]
    _sheet(wb, "Зоны внимания", ["Зона", "Объект", "Регион", "ОНПО", f"Топливо, {unit}",
                                 f"Прошлая неделя, {unit}", "Δ н/н, %", "Дней без продаж",
                                 "Негативных оценок", "Средняя оценка", "Главная категория негатива"],
           rows, [None, None, None, None, INT, INT, DELTA, INT, INT, "0.000", None],
           [20, 14, 28, 16, 16, 18, 10, 14, 14, 14, 28])

    charts = model["dynamics"]["charts"]
    if charts:
        header = ["Неделя"]
        for c in charts:
            for s in c["series"]:
                header.append(f"{c['title']} — {s['name']}")
        rows = []
        for i, label in enumerate(charts[0]["x"]):
            row = [label]
            for c in charts:
                for s in c["series"]:
                    row.append(s["values"][i])
            rows.append(row)
        _sheet(wb, "Динамика 8 недель", header, rows, [None] + [INT] * (len(header) - 1),
               [14] + [26] * (len(header) - 1))

    ws = wb.create_sheet("Паспорт")
    ws.column_dimensions["A"].width = 140
    for line in passport_lines(model):
        ws.append([line])
        ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    if model["pending"]:
        ws.append([", ".join(model["pending"]) + " — появятся в справке после дополнения витрины."])
    ws.append(["Конфиденциально. LUKOIL ОНПО · Инструмент АУП."])

    wb.properties.title = f"{model['title']} · {model['week']['iso']}"
    wb.save(path)
    return path
